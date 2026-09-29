"""
Módulo de serviço Excel - leitura e escrita de registos VIES.
"""

from pathlib import Path
from typing import Optional, Dict, List
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import os
import time


SHEET_NIF = "NIF"
COLUNA_NIF = "NIFS"
COLUNA_VALIDO = "Válido no VIES"
COLUNA_DATA = "Data consulta"


class ViesExcel:
    """Gerencia leitura e escrita de ficheiros Excel com validações VIES."""

    def __init__(self, caminho_ficheiro: str):
        self.caminho_ficheiro = Path(caminho_ficheiro)
        self.workbook = None
        self.ws = None
        self.colunas = None

    def abrir(self) -> None:
        """Abre o ficheiro Excel e localiza as colunas necessárias."""
        try:
            self.workbook = load_workbook(str(self.caminho_ficheiro))
        except Exception as exc:
            raise RuntimeError(f"Erro ao abrir o Excel: {exc}") from exc

        if SHEET_NIF not in self.workbook.sheetnames:
            self.workbook.close()
            raise RuntimeError(
                f"A sheet '{SHEET_NIF}' não foi encontrada."
            )

        self.ws = self.workbook[SHEET_NIF]
        self._localizar_colunas()

    def _localizar_colunas(self) -> None:
        """Localiza as colunas através dos cabeçalhos da linha 1."""
        self.colunas = {}

        for coluna in range(1, self.ws.max_column + 1):
            valor = self.ws.cell(row=1, column=coluna).value

            if valor is None:
                continue

            nome = str(valor).strip()
            self.colunas[nome] = coluna

        obrigatorias = [COLUNA_NIF, COLUNA_VALIDO, COLUNA_DATA]

        em_falta = [
            coluna
            for coluna in obrigatorias
            if coluna not in self.colunas
        ]

        if em_falta:
            self.workbook.close()
            raise RuntimeError(
                "Colunas obrigatórias não encontradas: "
                + ", ".join(em_falta)
            )

    def obter_registos(self) -> List[Dict]:
        """
        Lê todos os VATs da coluna NIFS para processamento.

        Processa SEMPRE todas as linhas, independentemente de terem sido
        processadas antes. O timestamp no Excel indica a última validação.
        """
        registos = []
        col_nif = self.colunas[COLUNA_NIF]

        for linha_excel in range(2, self.ws.max_row + 1):
            valor = self.ws.cell(row=linha_excel, column=col_nif).value

            if valor is None:
                continue

            vat = str(valor).strip()

            if not vat:
                continue

            registos.append(
                {
                    "linha": linha_excel,
                    "vat": vat,
                }
            )

        return registos

    def atualizar_resultado(
        self,
        linha: int,
        valido: bool,
        data_consulta: Optional[str],
    ) -> None:
        """Atualiza uma linha com o resultado da validação VIES."""
        col_valido = self.colunas[COLUNA_VALIDO]
        col_data = self.colunas[COLUNA_DATA]

        self.ws.cell(row=linha, column=col_valido).value = (
            "SIM" if valido else "NÃO"
        )

        self.ws.cell(row=linha, column=col_data).value = (
            data_consulta or ""
        )

    def atualizar_erro(self, linha: int) -> None:
        """Marca uma linha com erro."""
        col_valido = self.colunas[COLUNA_VALIDO]
        col_data = self.colunas[COLUNA_DATA]

        self.ws.cell(row=linha, column=col_valido).value = "ERRO"
        self.ws.cell(row=linha, column=col_data).value = ""

    def guardar(self) -> None:
        """Guarda o ficheiro Excel."""
        if self.workbook is None:
            return

        max_tentativas = 3
        tentativa = 0

        while tentativa < max_tentativas:
            try:
                self.workbook.save(str(self.caminho_ficheiro))
                return
            except PermissionError:
                tentativa += 1
                if tentativa < max_tentativas:
                    time.sleep(0.5)
                else:
                    raise RuntimeError(
                        "Não é possível gravar o ficheiro Excel. "
                        "Verifique se está aberto noutro programa."
                    )
            except Exception as exc:
                raise RuntimeError(f"Erro ao gravar o Excel: {exc}") from exc

    def fechar(self) -> None:
        """Fecha o ficheiro Excel."""
        if self.workbook is not None:
            self.workbook.close()
            self.workbook = None

    def __enter__(self):
        self.abrir()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.fechar()
