"""
Script auxiliar para criar um ficheiro Excel de teste.
Cria uma sheet 'NIF' com cabeçalhos e alguns VATs de teste.
"""

from pathlib import Path
from openpyxl import Workbook


def criar_teste_excel():
    """Cria um ficheiro Excel com dados de teste."""

    wb = Workbook()
    ws = wb.active
    ws.title = "NIF"

    # Cabeçalhos
    ws["A1"] = "NIFS"
    ws["B1"] = "Válido no VIES"
    ws["C1"] = "Data consulta"

    # Dados de teste
    testes = [
        ("PT504772694",),
        ("ESA28017895",),
        ("FRXY123456789",),
        ("INVALID",),
    ]

    for idx, (vat,) in enumerate(testes, start=2):
        ws[f"A{idx}"] = vat

    # Ajustar larguras
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 25

    # Guardar
    caminho = Path(__file__).parent / "teste_viex.xlsx"
    wb.save(str(caminho))

    print(f"Ficheiro de teste criado: {caminho}")
    print("Contém 4 VATs de teste na sheet 'NIF'")
    print("")


if __name__ == "__main__":
    criar_teste_excel()
