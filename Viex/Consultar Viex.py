from dataclasses import dataclass
from pathlib import Path
import os
import re
import requests
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from tkinter import Tk, filedialog


# ============================================================
# CONFIGURAÇÃO
# ============================================================

VIES_URL = (
    "https://ec.europa.eu/taxation_customs/vies/services/checkVatService"
)

LARGURA = 70

SHEET_NIF = "NIF"
COLUNA_NIF = "NIFS"
COLUNA_VALIDO = "Válido no VIES"
COLUNA_DATA = "Data consulta"


# ============================================================
# MODELO
# ============================================================

@dataclass
class ViesResult:
    vat: str
    country_code: str
    vat_number: str
    valid: bool
    name: str | None = None
    address: str | None = None
    request_date: str | None = None


# ============================================================
# UTILITÁRIOS DE APRESENTAÇÃO
# ============================================================

def limpar_ecra() -> None:
    os.system("cls" if os.name == "nt" else "clear")


def linha() -> None:
    print("=" * LARGURA)


def titulo() -> None:
    linha()
    print("VALIDADOR VIES".center(LARGURA))
    linha()


def pausar() -> None:
    input("\nPrima ENTER para continuar...")


# ============================================================
# TRATAMENTO DO VAT
# ============================================================

def normalizar_vat(vat: str) -> str:
    if not vat:
        raise ValueError("O VAT não pode estar vazio.")

    vat = str(vat).strip().upper()

    # Remove espaços, pontos, hífenes e outros separadores
    vat = re.sub(r"[^A-Z0-9]", "", vat)

    if len(vat) < 3:
        raise ValueError(
            "VAT inválido. Informe o código do país seguido do número."
        )

    return vat


def separar_vat(vat: str) -> tuple[str, str]:
    vat = normalizar_vat(vat)

    country_code = vat[:2]
    vat_number = vat[2:]

    if not country_code.isalpha():
        raise ValueError(
            "O VAT deve começar pelo código do país de 2 letras."
        )

    if not vat_number:
        raise ValueError("Número VAT não informado.")

    return country_code, vat_number


# ============================================================
# SERVIÇO VIES
# ============================================================

def validar_vies(
    vat: str,
    timeout: int = 15,
) -> ViesResult:

    country_code, vat_number = separar_vat(vat)

    soap_body = f"""<?xml version="1.0" encoding="UTF-8"?>
<soapenv:Envelope
    xmlns:soapenv="http://schemas.xmlsoap.org/soap/envelope/"
    xmlns:urn="urn:ec.europa.eu:taxud:vies:services:checkVat:types">

    <soapenv:Header/>

    <soapenv:Body>
        <urn:checkVat>
            <urn:countryCode>{country_code}</urn:countryCode>
            <urn:vatNumber>{vat_number}</urn:vatNumber>
        </urn:checkVat>
    </soapenv:Body>

</soapenv:Envelope>
"""

    headers = {
        "Content-Type": "text/xml; charset=utf-8",
        "SOAPAction": "",
    }

    try:
        response = requests.post(
            VIES_URL,
            data=soap_body.encode("utf-8"),
            headers=headers,
            timeout=timeout,
        )

        response.raise_for_status()

    except requests.Timeout as exc:
        raise RuntimeError(
            "O serviço VIES excedeu o tempo limite de resposta."
        ) from exc

    except requests.RequestException as exc:
        raise RuntimeError(
            f"Erro de comunicação com o serviço VIES: {exc}"
        ) from exc

    try:
        root = ET.fromstring(response.content)

    except ET.ParseError as exc:
        raise RuntimeError(
            "O VIES devolveu uma resposta XML inválida."
        ) from exc

    namespaces = {
        "soap": "http://schemas.xmlsoap.org/soap/envelope/",
        "ns": "urn:ec.europa.eu:taxud:vies:services:checkVat:types",
    }

    # --------------------------------------------------------
    # Verificar SOAP Fault
    # --------------------------------------------------------

    fault = root.find(".//soap:Fault", namespaces)

    if fault is not None:
        fault_code = fault.findtext("faultcode")
        fault_string = fault.findtext("faultstring")

        mensagem = (
            fault_string
            or fault_code
            or "Erro desconhecido."
        )

        raise RuntimeError(
            f"O serviço VIES devolveu um erro: {mensagem}"
        )

    # --------------------------------------------------------
    # Ler resposta
    # --------------------------------------------------------

    result = root.find(
        ".//ns:checkVatResponse",
        namespaces,
    )

    if result is None:
        raise RuntimeError(
            "Resposta inesperada do serviço VIES."
        )

    def get_text(field: str) -> str | None:
        element = result.find(
            f"ns:{field}",
            namespaces,
        )

        if element is None or element.text is None:
            return None

        value = element.text.strip()

        if value in {"", "---"}:
            return None

        return value

    valid_text = get_text("valid")

    return ViesResult(
        vat=f"{country_code}{vat_number}",
        country_code=get_text("countryCode") or country_code,
        vat_number=get_text("vatNumber") or vat_number,
        valid=(valid_text or "").lower() == "true",
        name=get_text("name"),
        address=get_text("address"),
        request_date=get_text("requestDate"),
    )


# ============================================================
# RESULTADO INDIVIDUAL
# ============================================================

def imprimir_resultado(resultado: ViesResult) -> None:

    print()
    linha()
    print("RESULTADO DA VALIDAÇÃO".center(LARGURA))
    linha()

    print(f"VAT:             {resultado.vat}")
    print(
        f"Válido no VIES:  "
        f"{'SIM' if resultado.valid else 'NÃO'}"
    )

    print(f"País:            {resultado.country_code}")
    print(f"Número:          {resultado.vat_number}")

    if resultado.valid:

        print(
            f"Nome:            "
            f"{resultado.name or 'NÃO DISPONIBILIZADO'}"
        )

        if resultado.address:
            linhas_morada = resultado.address.splitlines()

            print(
                f"Morada:          "
                f"{linhas_morada[0]}"
            )

            for parte in linhas_morada[1:]:
                print(
                    f"                 "
                    f"{parte}"
                )

        else:
            print(
                "Morada:          "
                "NÃO DISPONIBILIZADA"
            )

    else:

        print("Nome:            -")
        print("Morada:          -")

    print(
        f"Data consulta:   "
        f"{resultado.request_date or '-'}"
    )

    linha()


# ============================================================
# PESQUISA INDIVIDUAL
# ============================================================

def pesquisa_individual() -> None:

    while True:

        limpar_ecra()
        titulo()

        print()
        print("PESQUISA INDIVIDUAL")
        print("-" * LARGURA)
        print()

        vat = input("VAT: ").strip()

        try:

            resultado = validar_vies(vat)

            imprimir_resultado(resultado)

        except ValueError as exc:

            print()
            linha()
            print("ERRO")
            linha()
            print(str(exc))
            linha()

        except RuntimeError as exc:

            print()
            linha()
            print("ERRO NA CONSULTA VIES")
            linha()
            print(str(exc))
            linha()

        except KeyboardInterrupt:
            return

        while True:

            print()
            print("O que pretende fazer?")
            print()
            print("1 - Nova pesquisa")
            print("2 - Voltar ao menu principal")
            print("0 - Sair")
            print()

            opcao = input("Opção: ").strip()

            if opcao == "1":
                break

            if opcao == "2":
                return

            if opcao == "0":
                print()
                print("Validador VIES terminado.")
                raise SystemExit

            print()
            print("Opção inválida.")


# ============================================================
# SELEÇÃO DO EXCEL
# ============================================================

def selecionar_ficheiro_excel() -> str | None:
    """
    Abre uma janela em primeiro plano para selecionar
    o ficheiro Excel.
    """

    root = Tk()

    root.withdraw()

    # Forçar a janela de seleção para primeiro plano
    root.attributes("-topmost", True)

    root.update()

    ficheiro = filedialog.askopenfilename(
        parent=root,
        title="Selecionar ficheiro para validação VIES",
        filetypes=[
            ("Ficheiros Excel", "*.xlsx"),
            ("Todos os ficheiros", "*.*"),
        ],
    )

    root.destroy()

    if not ficheiro:
        return None

    return ficheiro


# ============================================================
# LOCALIZAR COLUNAS
# ============================================================

def localizar_colunas(ws) -> dict[str, int]:
    """
    Localiza as colunas através dos cabeçalhos da linha 1.
    """

    colunas = {}

    for coluna in range(1, ws.max_column + 1):

        valor = ws.cell(
            row=1,
            column=coluna,
        ).value

        if valor is None:
            continue

        nome = str(valor).strip()

        colunas[nome] = coluna

    obrigatorias = [
        COLUNA_NIF,
        COLUNA_VALIDO,
        COLUNA_DATA,
    ]

    em_falta = [
        coluna
        for coluna in obrigatorias
        if coluna not in colunas
    ]

    if em_falta:
        raise ValueError(
            "Colunas obrigatórias não encontradas: "
            + ", ".join(em_falta)
        )

    return colunas


# ============================================================
# PESQUISA EM MASSA
# ============================================================

def pesquisa_em_massa() -> None:

    limpar_ecra()
    titulo()

    print()
    print("PESQUISA EM MASSA")
    print("-" * LARGURA)
    print()
    print("Selecione o ficheiro Excel...")
    print()

    # --------------------------------------------------------
    # Selecionar ficheiro
    # --------------------------------------------------------

    ficheiro = selecionar_ficheiro_excel()

    if not ficheiro:

        print("Nenhum ficheiro selecionado.")
        pausar()
        return

    caminho = Path(ficheiro)

    print(f"Ficheiro: {caminho.name}")
    print()

    # --------------------------------------------------------
    # Abrir Excel
    # --------------------------------------------------------

    try:

        workbook = load_workbook(ficheiro)

    except Exception as exc:

        print(f"Erro ao abrir o Excel: {exc}")
        pausar()
        return

    # --------------------------------------------------------
    # Verificar sheet NIF
    # --------------------------------------------------------

    if SHEET_NIF not in workbook.sheetnames:

        print(
            f"ERRO: A sheet '{SHEET_NIF}' "
            "não foi encontrada."
        )

        workbook.close()
        pausar()
        return

    ws = workbook[SHEET_NIF]

    # --------------------------------------------------------
    # Localizar colunas
    # --------------------------------------------------------

    try:

        colunas = localizar_colunas(ws)

    except ValueError as exc:

        print(f"ERRO: {exc}")

        workbook.close()
        pausar()
        return

    col_nif = colunas[COLUNA_NIF]
    col_valido = colunas[COLUNA_VALIDO]
    col_data = colunas[COLUNA_DATA]

    # --------------------------------------------------------
    # Obter NIFs
    # --------------------------------------------------------

    registos = []

    for linha_excel in range(2, ws.max_row + 1):

        valor = ws.cell(
            row=linha_excel,
            column=col_nif,
        ).value

        if valor is None:
            continue

        vat = str(valor).strip()

        if not vat:
            continue

        registos.append(
            {
                "linha": linha_excel,
                "vat": vat,
            }
        )

    if not registos:

        print("Nenhum VAT encontrado na sheet NIF.")

        workbook.close()
        pausar()
        return

    # --------------------------------------------------------
    # Mostrar lista
    # --------------------------------------------------------

    total = len(registos)

    linha()

    print(
        f"VAT encontrados: {total}"
    )

    linha()
    print()

    for indice, registo in enumerate(
        registos,
        start=1,
    ):

        print(
            f"[ ] "
            f"{indice:>3}/{total}  "
            f"{registo['vat']}"
        )

    print()
    linha()
    print("A iniciar validação...")
    linha()
    print()

    # --------------------------------------------------------
    # Processar cada VAT
    # --------------------------------------------------------

    validos = 0
    invalidos = 0
    erros = 0

    for indice, registo in enumerate(
        registos,
        start=1,
    ):

        linha_excel = registo["linha"]
        vat = registo["vat"]

        try:

            resultado = validar_vies(vat)

            # -----------------------------------------------
            # Atualizar Excel
            # -----------------------------------------------

            ws.cell(
                row=linha_excel,
                column=col_valido,
            ).value = (
                "SIM"
                if resultado.valid
                else "NÃO"
            )

            ws.cell(
                row=linha_excel,
                column=col_data,
            ).value = (
                resultado.request_date or ""
            )

            # -----------------------------------------------
            # Guardar imediatamente
            # -----------------------------------------------

            workbook.save(ficheiro)

            # -----------------------------------------------
            # Estatísticas
            # -----------------------------------------------

            if resultado.valid:
                validos += 1
                estado = "VÁLIDO"

            else:
                invalidos += 1
                estado = "NÃO VÁLIDO"

            print(
                f"[✓] "
                f"{indice:>3}/{total}  "
                f"{vat:<20} "
                f"{estado}"
            )

        except Exception as exc:

            erros += 1

            # -----------------------------------------------
            # Registar erro no Excel
            # -----------------------------------------------

            ws.cell(
                row=linha_excel,
                column=col_valido,
            ).value = "ERRO"

            ws.cell(
                row=linha_excel,
                column=col_data,
            ).value = ""

            try:
                workbook.save(ficheiro)
            except Exception:
                pass

            print(
                f"[X] "
                f"{indice:>3}/{total}  "
                f"{vat:<20} "
                f"ERRO: {exc}"
            )

    # --------------------------------------------------------
    # Finalização
    # --------------------------------------------------------

    workbook.close()

    print()
    linha()
    print("VALIDAÇÃO CONCLUÍDA".center(LARGURA))
    linha()

    print()
    print(f"Total:        {total}")
    print(f"Válidos:      {validos}")
    print(f"Não válidos:  {invalidos}")
    print(f"Erros:        {erros}")

    print()
    print(f"Excel atualizado:")
    print(str(caminho))

    print()
    linha()

    pausar()


# ============================================================
# MENU PRINCIPAL
# ============================================================

def menu_principal() -> None:

    while True:

        limpar_ecra()
        titulo()

        print()
        print("Selecione o tipo de pesquisa:")
        print()
        print("1 - Pesquisa individual")
        print("2 - Pesquisa em massa")
        print("0 - Sair")
        print()

        opcao = input("Opção: ").strip()

        if opcao == "1":

            pesquisa_individual()

        elif opcao == "2":

            pesquisa_em_massa()

        elif opcao == "0":

            limpar_ecra()
            titulo()

            print()
            print("Validador VIES terminado.")
            print()

            break

        else:

            print()
            print("Opção inválida.")

            pausar()


# ============================================================
# MAIN
# ============================================================

def main() -> None:

    try:

        menu_principal()

    except KeyboardInterrupt:

        print()
        print()
        print("Operação cancelada.")


if __name__ == "__main__":
    main()