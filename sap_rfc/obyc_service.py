from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import traceback
import unicodedata
from pathlib import Path
from typing import Any

try:
    from pyrfc import Connection  # type: ignore
except Exception as exc:  # pragma: no cover - runtime guard
    Connection = None  # type: ignore[assignment]
    _PYRFC_IMPORT_ERROR = exc
else:
    _PYRFC_IMPORT_ERROR = None

from sap_rfc._rfc_common import (
    build_connection_params_for_env,
    classify_import_error,
    classify_rfc_error,
    find_project_root,
    format_exception,
    load_project_env,
    make_read_only_guard,
    resolve_target_env,
)

OBYC_ALLOWED_TABLES = ("T030", "T030K", "T030R", "T030B", "T030H")
OBYC_VALIDATION_PRIMARY_FIELDS = ("KTOPL", "KTOSL", "BWMOD", "KOMOK", "BKLAS")
OBYC_VALIDATION_REQUEST_FIELDS = ("KTOPL", "KTOSL", "BWMOD", "KOMOK", "BKLAS", "KONTS", "KONTH")
OBYC_VALIDATION_SAFE_FIELD_RE = re.compile(r"^[A-Z0-9_]{1,30}$")
OBYC_VALIDATION_FIELD_MAP = {
    "PLANO DE CONTAS": "KTOPL",
    "PLANO CONTAS": "KTOPL",
    "OPERACAO": "KTOSL",
    "OPERAÇÃO": "KTOSL",
    "COD AGRUP AVALIACAO": "BWMOD",
    "COD. AGRUP. AVALIACAO": "BWMOD",
    "COD AGRUP AVALIAÇÃO": "BWMOD",
    "CÓD.AGRUP.AVALIAÇÃO": "BWMOD",
    "MODIFICACAO CONTAS": "KOMOK",
    "MODIFICAÇÃO CONTAS": "KOMOK",
    "CLASSE DE AVALIACAO": "BKLAS",
    "CLASSE DE AVALIAÇÃO": "BKLAS",
    "CONTA DO RAZAO": "KONTS",
    "CONTA DO RAZÃO": "KONTS",
}


def _prepare_project_imports() -> None:
    project_root = os.getenv("SAP_SCRIPT_PROJECT_DIR", "").strip()
    if project_root and project_root not in sys.path:
        sys.path.insert(0, project_root)
    cockpit_dir = Path(project_root) / "sap_script_web_cockpit_v2" if project_root else None
    if cockpit_dir and cockpit_dir.exists():
        cockpit_dir_str = str(cockpit_dir)
        if cockpit_dir_str not in sys.path:
            sys.path.insert(0, cockpit_dir_str)


def _is_windows_runtime() -> bool:
    return os.name == "nt" or sys.platform.startswith("win")


def _bridge_python_executable() -> str | None:
    candidates = [
        os.getenv("SAP_FI_BRIDGE_PYTHON", "").strip(),
    ]
    if _is_windows_runtime():
        candidates.append(str((Path(__file__).resolve().parents[1] / ".venv-rfc" / "Scripts" / "python.exe").resolve()))
    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            return candidate
    return None


def _normalize_excel_preview_value(value: Any) -> str:
    if value is None:
        return ""
    try:
        from datetime import date, datetime

        if isinstance(value, (datetime, date)):
            return value.isoformat()
    except Exception:
        pass
    return str(value).strip()


def _load_excel_rows_for_validation(excel_path: str, sheet_name: str | None = None) -> dict[str, Any]:
    _prepare_project_imports()
    import openpyxl

    if not excel_path:
        raise RuntimeError("Caminho do Excel vazio.")

    suffix = Path(excel_path).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise RuntimeError("Formato de Excel não suportado. Use .xlsx ou .xlsm.")

    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if not workbook.sheetnames:
            raise RuntimeError("O ficheiro Excel não contém folhas visíveis.")

        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, ())

        display_headers: list[str] = []
        validation_headers: list[str] = []
        for index, value in enumerate(header_row, start=1):
            raw_header = _normalize_excel_preview_value(value)
            display_headers.append(raw_header if raw_header else f"COL{index}")
            normalized_header = _normalize_obyc_excel_header(raw_header)
            validation_headers.append(normalized_header if normalized_header else f"COL{index}")
        if not display_headers:
            display_headers = ["COL1"]
        if not validation_headers:
            validation_headers = ["COL1"]

        rows: list[dict[str, Any]] = []
        total_rows = 0
        for row_number, raw_row in enumerate(rows_iter, start=2):
            normalized_row: dict[str, Any] = {"_row_number": row_number}
            has_value = False
            for index, header in enumerate(display_headers):
                value = _normalize_excel_preview_value(raw_row[index]) if index < len(raw_row) else ""
                normalized_row[header] = value
                validation_header = validation_headers[index] if index < len(validation_headers) else header
                normalized_row[validation_header] = value
                if value:
                    has_value = True
            if not has_value:
                continue
            rows.append(normalized_row)
            total_rows += 1

        return {
            "ok": True,
            "file_name": Path(excel_path).name,
            "excel_path": excel_path,
            "sheet_name": sheet.title,
            "headers": display_headers,
            "validation_headers": validation_headers,
            "rows": rows,
            "row_count": total_rows,
        }
    finally:
        workbook.close()


def _normalize_obyc_excel_header(header: str) -> str:
    raw = str(header or "").strip()
    if not raw:
        return ""
    normalized = unicodedata.normalize("NFKD", raw)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.upper()
    normalized = re.sub(r"[^A-Z0-9]+", " ", normalized).strip()
    return OBYC_VALIDATION_FIELD_MAP.get(normalized, normalized.replace(" ", "_"))


def _obyc_validation_query_fields(row: dict[str, Any]) -> list[str]:
    candidate_fields: list[str] = []
    for field in OBYC_VALIDATION_PRIMARY_FIELDS:
        value = str(row.get(field) or "").strip()
        if value and field not in candidate_fields:
            candidate_fields.append(field)

    if len(candidate_fields) < 3:
        for field, value in row.items():
            if field == "_row_number":
                continue
            field_name = str(field or "").strip().upper()
            if not field_name or field_name in candidate_fields:
                continue
            if not OBYC_VALIDATION_SAFE_FIELD_RE.match(field_name):
                continue
            if str(value or "").strip():
                candidate_fields.append(field_name)
            if len(candidate_fields) >= 3:
                break

    return candidate_fields[:5]


def _obyc_validation_comparable_fields(excel_row: dict[str, Any], rfc_row: dict[str, Any]) -> list[str]:
    fields: list[str] = []
    for field in excel_row.keys():
        if field == "_row_number":
            continue
        if field in rfc_row and field not in fields:
            fields.append(field)
        if field == "KONTS" and "KONTH" in rfc_row and "KONTH" not in fields:
            fields.append("KONTH")
        if field == "KONTH" and "KONTS" in rfc_row and "KONTS" not in fields:
            fields.append("KONTS")
    return fields


def _obyc_validation_key_fields(excel_row: dict[str, Any], rfc_row: dict[str, Any]) -> list[str]:
    fields: list[str] = []
    for field in OBYC_VALIDATION_PRIMARY_FIELDS:
        if field in excel_row and field in rfc_row:
            fields.append(field)
    return fields


def _obyc_validation_optional_fields(excel_row: dict[str, Any]) -> list[str]:
    fields: list[str] = []
    primary_fields = set(OBYC_VALIDATION_PRIMARY_FIELDS)
    for field in excel_row.keys():
        if field == "_row_number" or field in primary_fields:
            continue
        field_name = str(field or "").strip().upper()
        if not field_name or field_name in primary_fields:
            continue
        if not OBYC_VALIDATION_SAFE_FIELD_RE.match(field_name):
            continue
        fields.append(field)
    return fields


def _obyc_row_value(row: dict[str, Any], field: str) -> str:
    if field in {"KONTS", "KONTH"}:
        first = str(row.get("KONTS") or "").strip()
        if first:
            return first.lstrip("0") or "0" if first.isdigit() else first
        second = str(row.get("KONTH") or "").strip()
        if second:
            return second.lstrip("0") or "0" if second.isdigit() else second
        return ""
    return str(row.get(field) or "").strip()


def _obyc_validation_filter_value(field: str, value: Any) -> str:
    text = str(value or "").strip()
    if field in {"KONTS", "KONTH"} and text.isdigit():
        return text.zfill(10)
    return text


def _obyc_connection_params(environment: str) -> dict[str, str]:
    prefix = f"SAP_{environment}_"
    required = ["USER", "PASSWD", "ASHOST", "SYSNR", "CLIENT"]
    missing = [f"{prefix}{key}" for key in required if not os.getenv(f"{prefix}{key}")]
    if missing:
        raise RuntimeError(f"Missing SAP environment variables: {', '.join(missing)}")

    return {
        "user": os.environ[f"{prefix}USER"],
        "passwd": os.environ[f"{prefix}PASSWD"],
        "ashost": os.environ[f"{prefix}ASHOST"],
        "sysnr": os.environ[f"{prefix}SYSNR"],
        "client": os.environ[f"{prefix}CLIENT"],
        "lang": os.getenv(f"{prefix}LANG", "EN"),
    }


def _is_obyc_no_data_error(exc: BaseException) -> bool:
    text = f"{exc.__class__.__name__} {getattr(exc, 'key', '')} {getattr(exc, 'message', '')} {exc}".upper()
    return "TABLE_WITHOUT_DATA" in text or "RFC_TABLE_WITHOUT_DATA" in text


def _obyc_read_table_core(environment: str, table: str, filters: list[dict[str, str]], fields: list[str]) -> list[dict[str, Any]]:
    _prepare_project_imports()

    try:
        from pyrfc import Connection  # type: ignore
        from sap_agent.safety import SafetyGuard
    except Exception as exc:
        raise RuntimeError(f"Não foi possível importar o cliente RFC read-only: {exc}") from exc

    conn = None
    try:
        safety_guard = SafetyGuard.build(
            allow_write_operations=False,
            allowed_functions=("RFC_PING", "RFC_READ_TABLE"),
            allowed_tables=OBYC_ALLOWED_TABLES,
        )
        safety_guard.assert_function_allowed("RFC_READ_TABLE")
        safety_guard.assert_table_allowed(table)
        conn = Connection(**_obyc_connection_params(environment))
        fields_payload = [{"FIELDNAME": field} for field in fields]
        options_payload = []
        for index, item in enumerate(filters):
            prefix = "AND " if index else ""
            value = str(item["value"]).replace("'", "''")
            options_payload.append(f"{prefix}{item['field']} = '{value}'")
        result = conn.call(
            "RFC_READ_TABLE",
            QUERY_TABLE=table,
            DELIMITER="|",
            FIELDS=fields_payload,
            OPTIONS=[{"TEXT": option} for option in options_payload],
            ROWCOUNT=50,
        )
        sap_fields = [
            str(entry.get("FIELDNAME") or "").strip()
            for entry in result.get("FIELDS", [])
            if str(entry.get("FIELDNAME") or "").strip()
        ]
        rows: list[dict[str, Any]] = []
        for row in result.get("DATA", []):
            values = str(row.get("WA", "")).split("|")
            rows.append(
                {
                    field: (values[index].strip() if index < len(values) else "")
                    for index, field in enumerate(sap_fields)
                }
            )
        return rows
    except Exception as exc:
        if _is_obyc_no_data_error(exc):
            return []
        raise RuntimeError(f"Falha na consulta read-only OBYC via RFC: {exc}") from exc
    finally:
        try:
            if conn is not None:
                conn.close()
        except Exception:
            pass


def _obyc_read_table_via_bridge(environment: str, table: str, filters: list[dict[str, str]], fields: list[str]) -> list[dict[str, Any]]:
    python_exe = _bridge_python_executable()
    if not python_exe:
        runtime = "Windows" if _is_windows_runtime() else "WSL/Linux"
        raise RuntimeError(
            f"Execução OBYC via bridge indisponível neste runtime ({runtime}). "
            "Configure SAP_FI_BRIDGE_PYTHON com um Python que tenha PyRFC instalado."
        )

    payload = {
        "environment": environment,
        "table": table,
        "filters": filters,
        "fields": fields,
    }
    bridge_code = (
        "import json, os, sys, traceback\n"
        "from pathlib import Path\n"
        "project_dir = os.environ.get('SAP_SCRIPT_PROJECT_DIR', '').strip()\n"
        "if project_dir:\n"
        "    project_path = Path(project_dir)\n"
        "    if str(project_path) not in sys.path:\n"
        "        sys.path.insert(0, str(project_path))\n"
        "payload = json.loads(sys.stdin.read() or '{}')\n"
        "from sap_rfc.obyc_service import _obyc_read_table_core\n"
        "try:\n"
        "    rows = _obyc_read_table_core(payload['environment'], payload['table'], payload['filters'], payload['fields'])\n"
        "    print(json.dumps({'ok': True, 'rows': rows}, ensure_ascii=False))\n"
        "except Exception as exc:\n"
        "    print(json.dumps({'ok': False, 'message': str(exc), 'traceback': traceback.format_exc()}, ensure_ascii=False))\n"
        "    raise SystemExit(1)\n"
    )
    proc = subprocess.run(
        [python_exe, "-c", bridge_code],
        input=json.dumps(payload, ensure_ascii=False),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        cwd=str(Path(__file__).resolve().parent.parent),
    )

    stdout = (proc.stdout or "").strip()
    if not stdout:
        raise RuntimeError(f"Bridge OBYC devolveu saída vazia.\nSTDERR: {proc.stderr}")

    try:
        data = json.loads(stdout.splitlines()[-1])
    except Exception as exc:
        raise RuntimeError(f"Bridge OBYC devolveu JSON inválido.\nSTDOUT: {proc.stdout}\nSTDERR: {proc.stderr}") from exc

    if not data.get("ok"):
        raise RuntimeError(str(data.get("message") or "Falha no bridge OBYC."))
    rows = data.get("rows")
    if not isinstance(rows, list):
        raise RuntimeError("Bridge OBYC devolveu linhas inválidas.")
    return [row for row in rows if isinstance(row, dict)]


def read_obyc_table(environment: str, table: str, filters: list[dict[str, str]], fields: list[str]) -> list[dict[str, Any]]:
    try:
        return _obyc_read_table_core(environment, table, filters, fields)
    except Exception as exc:
        if Connection is None:
            return _obyc_read_table_via_bridge(environment, table, filters, fields)
        raise RuntimeError(str(exc)) from exc


def _build_workbook_from_preview_data(preview_data: dict[str, Any]) -> dict[str, Any]:
    return {
        "file_name": str(preview_data.get("file_name") or "Excel da OBYC").strip() or "Excel da OBYC",
        "sheet_name": str(preview_data.get("sheet_name") or "Folha principal").strip() or "Folha principal",
        "headers": list(preview_data.get("headers") or []),
        "rows": [row for row in (preview_data.get("rows") or []) if isinstance(row, dict)],
        "row_count": int(preview_data.get("row_count") or len(preview_data.get("rows") or [])),
    }


def _load_workbook_from_preview_job(preview_job_id: str) -> dict[str, Any]:
    _prepare_project_imports()
    from web_api.store import get_job

    preview_job = get_job(preview_job_id)
    if not preview_job:
        raise RuntimeError("Job de pré-visualização do Excel não encontrado.")
    if preview_job.get("task") != "obyc_excel_preview":
        raise RuntimeError("O job informado não pertence ao fluxo de pré-visualização da OBYC.")

    raw_status = preview_job.get("status")
    try:
        preview_result = json.loads(raw_status) if isinstance(raw_status, str) else raw_status
    except Exception as exc:
        raise RuntimeError(f"Não foi possível ler o resultado do Excel OBYC: {exc}") from exc

    if not isinstance(preview_result, dict):
        raise RuntimeError("Resultado da pré-visualização do Excel OBYC inválido.")

    file_name = str(preview_result.get("file_name") or "").strip()
    sheet_name = str(preview_result.get("sheet_name") or "").strip()
    workbook_rows = preview_result.get("all_rows")
    if isinstance(workbook_rows, list):
        return {
            "file_name": file_name or "Excel da OBYC",
            "sheet_name": sheet_name or "Folha principal",
            "headers": list(preview_result.get("validation_headers") or preview_result.get("headers") or []),
            "rows": [row for row in workbook_rows if isinstance(row, dict)],
            "row_count": int(preview_result.get("validation_row_count") or preview_result.get("row_count") or len(workbook_rows)),
        }

    excel_path = str(preview_result.get("excel_path") or "").strip()
    if not excel_path:
        raise RuntimeError("O caminho do ficheiro Excel OBYC não está disponível.")
    return _load_excel_rows_for_validation(excel_path, sheet_name or None)


def validate_obyc_excel(params: dict[str, Any]) -> tuple[str, str]:
    preview_job_id = str(params.get("preview_job_id") or "").strip()
    preview_data = params.get("preview_data")
    environment = str(params.get("system") or "DEV").strip().upper() or "DEV"
    table = str(params.get("table") or "T030").strip().upper() or "T030"

    if table not in OBYC_ALLOWED_TABLES:
        raise RuntimeError(
            f"Tabela OBYC não permitida: {table}. Tabelas permitidas: {', '.join(OBYC_ALLOWED_TABLES)}."
        )

    if isinstance(preview_data, dict) and preview_data.get("rows"):
        workbook = _build_workbook_from_preview_data(preview_data)
    else:
        if not preview_job_id:
            raise RuntimeError("Job de pré-visualização do Excel não informado.")
        workbook = _load_workbook_from_preview_job(preview_job_id)

    excel_rows = [row for row in workbook["rows"] if isinstance(row, dict)]

    validated_rows = 0
    matched_rows = 0
    missing_rows = 0
    mismatched_rows = 0
    optional_rows = 0
    skipped_rows = 0
    issues: list[dict[str, Any]] = []
    query_cache: dict[tuple[tuple[str, str], ...], list[dict[str, Any]]] = {}

    for excel_row in excel_rows:
        row_number = int(excel_row.get("_row_number") or 0)
        query_fields = _obyc_validation_query_fields(excel_row)
        filters = [
            {"field": field, "value": _obyc_validation_filter_value(field, excel_row.get(field))}
            for field in query_fields
            if str(excel_row.get(field) or "").strip()
        ]
        if not filters:
            skipped_rows += 1
            issues.append(
                {
                    "row_number": row_number,
                    "reason": "sem_chaves",
                    "message": "A linha não contém chaves suficientes para validação.",
                }
            )
            continue

        validated_rows += 1
        cache_key = tuple((item["field"], item["value"]) for item in filters)
        if cache_key not in query_cache:
            query_cache[cache_key] = read_obyc_table(
                environment,
                table,
                filters,
                [
                    field
                    for field in excel_row.keys()
                    if field in OBYC_VALIDATION_REQUEST_FIELDS
                ],
            )
        sap_rows = query_cache[cache_key]

        if not sap_rows:
            missing_rows += 1
            issues.append(
                {
                    "row_number": row_number,
                    "reason": "nao_encontrado",
                    "filters": filters,
                    "message": f"Nenhum registo encontrado em {table} para as chaves informadas.",
                }
            )
            continue

        row_matched = False
        exact_optional_match = False
        best_optional_differences: list[dict[str, Any]] = []
        optional_fields = _obyc_validation_optional_fields(excel_row)
        for sap_row in sap_rows:
            key_fields = _obyc_validation_key_fields(excel_row, sap_row)
            key_differences = []
            for field in key_fields:
                excel_value = _obyc_row_value(excel_row, field)
                sap_value = _obyc_row_value(sap_row, field)
                if excel_value != sap_value:
                    key_differences.append(
                        {
                            "field": field,
                            "excel": excel_value,
                            "sap": sap_value,
                        }
                    )
            if key_differences:
                continue

            row_matched = True
            optional_differences: list[dict[str, Any]] = []
            for field in optional_fields:
                excel_value = _obyc_row_value(excel_row, field)
                sap_value = _obyc_row_value(sap_row, field)
                if excel_value != sap_value:
                    optional_differences.append(
                        {
                            "field": field,
                            "excel": excel_value,
                            "sap": sap_value,
                        }
                    )

            if not optional_differences:
                exact_optional_match = True
                break
            if not best_optional_differences or len(optional_differences) < len(best_optional_differences):
                best_optional_differences = optional_differences

        if row_matched:
            matched_rows += 1
            if not exact_optional_match and best_optional_differences:
                optional_rows += 1
                issues.append(
                    {
                        "row_number": row_number,
                        "reason": "comparacao_opcional",
                        "filters": filters,
                        "message": "Campos opcionais do Excel diferem do registo SAP, mas as chaves principais coincidem.",
                        "differences": best_optional_differences[:12],
                    }
                )
            continue

        mismatched_rows += 1
        issues.append(
            {
                "row_number": row_number,
                "reason": "chave_existente_conta_diferente",
                "filters": filters,
                "differences": [],
                "message": f"Chave existente em {table} com contas diferentes.",
            }
        )

    message = (
        f"Validação do Excel concluída em {table} ({environment}). "
        f"Linhas verificadas: {validated_rows}. "
        f"Coincidentes: {matched_rows}. "
        f"Comparações opcionais com diferenças: {optional_rows}. "
        f"Chave existente com contas diferentes: {mismatched_rows}. "
        f"Sem correspondência: {missing_rows}. "
        f"Ignoradas: {skipped_rows}."
    )

    result = {
        "ok": True,
        "status": "VALIDATION_READY",
        "system": environment,
        "table": table,
        "file_name": workbook["file_name"],
        "sheet_name": workbook["sheet_name"],
        "headers": workbook["headers"],
        "total_rows": workbook["row_count"],
        "validated_rows": validated_rows,
        "matched_rows": matched_rows,
        "optional_rows": optional_rows,
        "missing_rows": missing_rows,
        "mismatched_rows": mismatched_rows,
        "skipped_rows": skipped_rows,
        "issues": issues[:20],
        "message": message,
    }
    log = "\n".join(
        [
            "Validação read-only do Excel OBYC executada via RFC.",
            f"Ambiente: {environment}",
            f"Tabela: {table}",
            f"Ficheiro: {workbook['file_name']}",
            f"Linhas verificadas: {validated_rows}",
            f"Coincidentes: {matched_rows}",
            f"Comparações opcionais com diferenças: {optional_rows}",
            f"Chave existente com contas diferentes: {mismatched_rows}",
            f"Sem correspondência: {missing_rows}",
            f"Ignoradas: {skipped_rows}",
        ]
    )
    return json.dumps(result, ensure_ascii=False), log
