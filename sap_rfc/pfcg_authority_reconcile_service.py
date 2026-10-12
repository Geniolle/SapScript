# -*- coding: utf-8 -*-
"""
sap_rfc/pfcg_authority_reconcile_service.py
======================================================================
Serviço Oficial de Reconciliação da Sheet PFCG_AUTHORITY com as sheets:
  - PFCG_CREATE
  - PFCG_COMPOSTA
======================================================================
Regras Oficiais:
1. PFCG_AUTHORITY.AGR_NAME_COMPOSTA preenchido -> TIPO_ESPERADO = COMPOSITE
   Relação (AGR_NAME_COMPOSTA, AGR_NAME) deve existir na PFCG_COMPOSTA quando houver AGR_NAME.
2. PFCG_AUTHORITY.AGR_NAME preenchido -> TIPO_ESPERADO = SINGLE
   Deve existir na PFCG_CREATE.
3. Se AGR_NAME_COMPOSTA preenchido e AGR_NAME vazio:
   - Classificar como AUTHORITY_COMPOSITE_SEM_CHILD.
   - Continua sendo COMPOSITE, mas não cria relação sem child em PFCG_COMPOSTA.
4. Se AGR_NAME_COMPOSTA vazio e AGR_NAME preenchido:
   - Single avulsa, deve existir em PFCG_CREATE, nenhuma relação em PFCG_COMPOSTA.
5. Se a mesma role aparecer como AGR_NAME_COMPOSTA e AGR_NAME:
   - Classificar como CONFLITO_TIPO_AUTHORITY (não resolver silenciosamente).
6. Idempotência e deduplicação:
   - PFCG_CREATE: chave AGR_NAME (trim + upper). Preserva descrição existente se já houver.
   - PFCG_COMPOSTA: chave (AGR_NAME_COMPOSTA, AGR_NAME) (trim + upper).
7. Modos:
   - dry_run=True: SAP_WRITE=0, EXCEL_WRITE=0 (somente análise e delta).
   - dry_run=False: backup timestampado obrigatório em output/ antes da escrita.
======================================================================
"""

from __future__ import annotations

import os
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple


def _norm_col(val: Any) -> str:
    """Normaliza nome de coluna para maiúsculas sem acentos."""
    if val is None:
        return ""
    return unicodedata.normalize("NFKD", str(val)).encode("ASCII", "ignore").decode("utf-8").strip().upper()


def _norm_val(val: Any) -> str:
    """Normaliza valor de texto/role."""
    if val is None:
        return ""
    s = str(val).strip().upper()
    return "" if s in ("NAN", "NONE", "-", "") else s


def analisar_reconciliacao_pfcg_authority(caminho_excel: str) -> Dict[str, Any]:
    """
    Analisa a folha PFCG_AUTHORITY e calcula o delta em relação a PFCG_CREATE e PFCG_COMPOSTA.
    Identificação dinâmica de colunas por cabeçalho, sem depender de posições fixas.
    Totalmente READ-ONLY.
    """
    import openpyxl

    wb = openpyxl.load_workbook(caminho_excel, read_only=True, data_only=True)
    sheet_names = wb.sheetnames

    # 1. Mapear PFCG_CREATE existente
    existing_create: Dict[str, Dict[str, Any]] = {}
    if "PFCG_CREATE" in sheet_names:
        ws_cr = wb["PFCG_CREATE"]
        header_cr = [_norm_col(c) for c in next(ws_cr.iter_rows(max_row=1, values_only=True), ())]
        col_cr_agr = next((i for i, c in enumerate(header_cr) if c in ("AGR_NAME", "ROLE")), -1)
        col_cr_txt = next((i for i, c in enumerate(header_cr) if c in ("TEXT", "DESCRICAO")), -1)
        col_cr_tc = next((i for i, c in enumerate(header_cr) if c in ("TCODE", "TRANSACTION")), -1)

        for row in ws_cr.iter_rows(min_row=2, values_only=True):
            if col_cr_agr >= 0 and col_cr_agr < len(row):
                role = _norm_val(row[col_cr_agr])
                if role:
                    txt = str(row[col_cr_txt]).strip() if col_cr_txt >= 0 and col_cr_txt < len(row) and row[col_cr_txt] is not None else ""
                    tc = str(row[col_cr_tc]).strip() if col_cr_tc >= 0 and col_cr_tc < len(row) and row[col_cr_tc] is not None else ""
                    existing_create[role] = {"text": txt, "tcode": tc}

    # 2. Mapear PFCG_COMPOSTA existente
    existing_composta_pairs: Set[Tuple[str, str]] = set()
    if "PFCG_COMPOSTA" in sheet_names:
        ws_cp = wb["PFCG_COMPOSTA"]
        header_cp = [_norm_col(c) for c in next(ws_cp.iter_rows(max_row=1, values_only=True), ())]
        col_cp_p = next((i for i, c in enumerate(header_cp) if c in ("AGR_NAME_COMPOSTA", "COMPOSITEROLE")), -1)
        col_cp_c = next((i for i, c in enumerate(header_cp) if c in ("AGR_NAME", "SINGLEROLE")), -1)

        for row in ws_cp.iter_rows(min_row=2, values_only=True):
            p = _norm_val(row[col_cp_p]) if col_cp_p >= 0 and col_cp_p < len(row) else ""
            c = _norm_val(row[col_cp_c]) if col_cp_c >= 0 and col_cp_c < len(row) else ""
            if p and c:
                existing_composta_pairs.add((p, c))

    # 3. Analisar PFCG_AUTHORITY
    linhas_analisadas = 0
    composites_encontradas: Set[str] = set()
    singles_encontradas: Set[str] = set()
    composite_sem_child_count = 0
    duplicados_ignorados = 0

    novas_create_list: List[Dict[str, str]] = []
    novas_composta_list: List[Dict[str, str]] = []

    seen_create_roles: Set[str] = set()
    seen_composta_pairs: Set[Tuple[str, str]] = set()

    # Mapeamento detalhado por role para auditoria de Business Partner e conflitos
    roles_origem: Dict[str, Dict[str, Any]] = {}

    if "PFCG_AUTHORITY" in sheet_names:
        ws_au = wb["PFCG_AUTHORITY"]
        header_au = [_norm_col(c) for c in next(ws_au.iter_rows(max_row=1, values_only=True), ())]

        col_au_comp = next((i for i, c in enumerate(header_au) if c in ("AGR_NAME_COMPOSTA", "COMPOSITEROLE")), -1)
        col_au_ind = next((i for i, c in enumerate(header_au) if c in ("AGR_NAME", "SINGLEROLE")), -1)

        # Tratar colunas TEXT (pode haver TEXT de composta e TEXT de individual)
        text_indices = [i for i, c in enumerate(header_au) if c in ("TEXT", "DESCRICAO")]
        col_au_txt_comp = text_indices[0] if len(text_indices) > 0 else -1
        col_au_txt_ind = text_indices[1] if len(text_indices) > 1 else (text_indices[0] if text_indices else -1)

        for row in ws_au.iter_rows(min_row=2, values_only=True):
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            linhas_analisadas += 1
            p = _norm_val(row[col_au_comp]) if col_au_comp >= 0 and col_au_comp < len(row) else ""
            c = _norm_val(row[col_au_ind]) if col_au_ind >= 0 and col_au_ind < len(row) else ""

            txt_comp = str(row[col_au_txt_comp]).strip() if col_au_txt_comp >= 0 and col_au_txt_comp < len(row) and row[col_au_txt_comp] is not None else ""
            txt_ind = str(row[col_au_txt_ind]).strip() if col_au_txt_ind >= 0 and col_au_txt_ind < len(row) and row[col_au_txt_ind] is not None else ""

            if p:
                composites_encontradas.add(p)
                if p not in roles_origem:
                    roles_origem[p] = {"aparece_composta": True, "aparece_single": False, "coluna": "AGR_NAME_COMPOSTA"}
                else:
                    roles_origem[p]["aparece_composta"] = True

            if c:
                singles_encontradas.add(c)
                if c not in roles_origem:
                    roles_origem[c] = {"aparece_composta": False, "aparece_single": True, "coluna": "AGR_NAME"}
                else:
                    roles_origem[c]["aparece_single"] = True

            # Caso 7: AGR_NAME_COMPOSTA preenchido e AGR_NAME vazio
            if p and not c:
                composite_sem_child_count += 1

            # Caso 5: Reconciliação PFCG_CREATE para todo AGR_NAME válido
            if c:
                if c not in existing_create and c not in seen_create_roles:
                    novas_create_list.append({
                        "role": c,
                        "text": txt_ind or c,
                        "tcode": ""
                    })
                    seen_create_roles.add(c)
                elif c in seen_create_roles:
                    duplicados_ignorados += 1

            # Caso 4: Reconciliação PFCG_COMPOSTA (p != vazio e c != vazio)
            if p and c:
                pair = (p, c)
                if pair not in seen_composta_pairs:
                    seen_composta_pairs.add(pair)
                    if pair not in existing_composta_pairs:
                        novas_composta_list.append({
                            "parent": p,
                            "child": c,
                            "text_parent": txt_comp or p
                        })
                else:
                    duplicados_ignorados += 1

    wb.close()

    # Caso 17: Validar Conflitos (mesma role aparece em AGR_NAME_COMPOSTA e AGR_NAME)
    conflitos_lista = []
    for r_nome, info in sorted(roles_origem.items()):
        if info["aparece_composta"] and info["aparece_single"]:
            conflitos_lista.append({
                "role": r_nome,
                "aparece_como_composite": True,
                "aparece_como_single": True,
                "status": "CONFLITO_TIPO_AUTHORITY"
            })

    # Classificação específica das 3 roles Business Partner
    bp_roles = ["Z_BR_BUSINESS_PARTNER", "Z_BR_BUSINESS_PARTNER_GL", "Z_BR_TYPE_BP_GERAL"]
    bp_classificacao = {}
    for bp in bp_roles:
        info = roles_origem.get(bp)
        if info:
            if info["aparece_composta"] and not info["aparece_single"]:
                tipo_esp = "COMPOSITE"
                col = "AGR_NAME_COMPOSTA"
            elif info["aparece_single"] and not info["aparece_composta"]:
                tipo_esp = "SINGLE"
                col = "AGR_NAME"
            else:
                tipo_esp = "CONFLITO_TIPO_AUTHORITY"
                col = "AMBAS"
            bp_classificacao[bp] = {
                "role": bp,
                "origem": "PFCG_AUTHORITY",
                "coluna": col,
                "tipo_esperado_resultante": tipo_esp
            }
        else:
            bp_classificacao[bp] = {
                "role": bp,
                "origem": "NAO_ENCONTRADA_EM_AUTHORITY",
                "coluna": "N/A",
                "tipo_esperado_resultante": "UNKNOWN"
            }

    ja_existentes_create = len(singles_encontradas) - len(novas_create_list)
    # Pares de PFCG_AUTHORITY que já existem na PFCG_COMPOSTA
    ja_existentes_composta = len(seen_composta_pairs) - len(novas_composta_list)

    return {
        "linhas_analisadas": linhas_analisadas,
        "composites_distintas": sorted(list(composites_encontradas)),
        "composites_distintas_count": len(composites_encontradas),
        "singles_distintas": sorted(list(singles_encontradas)),
        "singles_distintas_count": len(singles_encontradas),
        "create_ja_existentes": ja_existentes_create,
        "create_novas_necessarias": novas_create_list,
        "create_novas_count": len(novas_create_list),
        "composta_ja_existentes": ja_existentes_composta,
        "composta_novas_necessarias": novas_composta_list,
        "composta_novas_count": len(novas_composta_list),
        "composite_sem_child_count": composite_sem_child_count,
        "duplicados_ignorados": duplicados_ignorados,
        "conflitos": conflitos_lista,
        "conflitos_count": len(conflitos_lista),
        "business_partner_classificacao": bp_classificacao,
        "roles_origem": roles_origem,
    }


def gerar_backup_timestampado(caminho_excel: str) -> Tuple[bool, str]:
    """
    Cria uma cópia de segurança timestampada na pasta output/ antes de qualquer gravação no Excel.
    Retorna (sucesso: bool, caminho_backup: str).
    """
    try:
        p_orig = Path(caminho_excel).resolve()
        dir_output = p_orig.parent / "output"
        dir_output.mkdir(parents=True, exist_ok=True)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        nome_backup = f"S4H_Perfis_v1_backup_{ts}.xlsx"
        p_dest = dir_output / nome_backup

        shutil.copy2(str(p_orig), str(p_dest))
        if p_dest.exists() and p_dest.stat().st_size > 0:
            return True, str(p_dest)
        return False, "Ficheiro de backup não foi criado ou tamanho zero."
    except Exception as exc:
        return False, str(exc)


def reconciliar_pfcg_authority_com_catalogos(
    caminho_excel: str,
    dry_run: bool = True,
    imprimir_log: bool = True,
    debug: bool = False
) -> Dict[str, Any]:
    """
    Executa a reconciliação entre PFCG_AUTHORITY e as folhas PFCG_CREATE e PFCG_COMPOSTA.

    Parâmetros:
      caminho_excel: Caminho absoluto ou relativo para o ficheiro mestre Excel.
      dry_run: Se True, SAP_WRITE=0 e EXCEL_WRITE=0 (somente análise e relatório delta).
               Se False, cria backup obrigatório e grava novas linhas nas folhas.
      imprimir_log: Se True, imprime o relatório estruturado no terminal.
      debug: Se True, imprime o relatório detalhado completo de diagnóstico.

    Retorna:
      Dicionário com o resultado completo da reconciliação.
    """
    if not debug:
        debug = os.environ.get("SAP_DEBUG", "").strip() in ("1", "true", "TRUE", "True")
    analise = analisar_reconciliacao_pfcg_authority(caminho_excel)

    if imprimir_log:
        linhas_totais = analise['linhas_analisadas']
        c_count = analise['composites_distintas_count']
        s_count = analise['singles_distintas_count']
        cr_exist = analise['create_ja_existentes']
        cr_novas = analise['create_novas_count']
        cr_total = cr_exist + cr_novas
        cp_exist = analise['composta_ja_existentes']
        cp_novas = analise['composta_novas_count']
        cp_total = cp_exist + cp_novas
        conflitos = analise['conflitos_count']

        tem_pendencia = (cr_novas > 0 or cp_novas > 0 or conflitos > 0)

        if debug:
            print("\n" + "=" * 78)
            print("  RECONCILIAÇÃO PFCG_AUTHORITY (DEBUG)")
            print("=" * 78)
            print(f"  Linhas analisadas                  : {linhas_totais}")
            print(f"  Composite distintas encontradas   : {c_count}")
            print(f"  Singles distintas encontradas     : {s_count}")
            print("\n  PFCG_CREATE:")
            print(f"    Já existentes                   : {cr_exist}")
            print(f"    Novas necessárias               : {cr_novas}")
            print("\n  PFCG_COMPOSTA:")
            print(f"    Relações já existentes          : {cp_exist}")
            print(f"    Novas relações necessárias      : {cp_novas}")
            print(f"\n  Authority sem child               : {analise['composite_sem_child_count']}")
            print(f"  Duplicados ignorados              : {analise['duplicados_ignorados']}")
            print(f"  Conflitos de tipo em Authority     : {conflitos}")
            print("=" * 78)

        if not tem_pendencia:
            from sap_rfc.status_logger import print_status_success
            print_status_success(f"  ✓ [5/8] PFCG_AUTHORITY: {linhas_totais} linhas | {c_count}C | {s_count}S | CREATE {cr_exist}/{cr_total} | COMPOSTA {cp_exist}/{cp_total} | conflitos {conflitos}")
        else:
            from sap_rfc.status_logger import print_status_warning
            txt_cr = f"{cr_exist}/{cr_total} (+{cr_novas})" if cr_novas > 0 else f"{cr_exist}/{cr_total}"
            txt_cp = f"{cp_exist}/{cp_total} (+{cp_novas})" if cp_novas > 0 else f"{cp_exist}/{cp_total}"
            print_status_warning(f"  ⚠️ [5/8] PFCG_AUTHORITY: {linhas_totais} linhas | {c_count}C | {s_count}S | CREATE {txt_cr} | COMPOSTA {txt_cp} | conflitos {conflitos}")

            if cr_novas > 0:
                print("  CREATE ausentes:")
                for item in analise["create_novas_necessarias"]:
                    print(f"    - {item['role']}")

            if cp_novas > 0:
                print("  COMPOSTA ausente:")
                for item in analise["composta_novas_necessarias"]:
                    print(f"    - {item['parent']} -> {item['child']}")

            if conflitos > 0:
                print("  Conflitos:")
                for conf in analise["conflitos"]:
                    print(f"    - {conf['role']}")

    resultado = {
        "ok": True,
        "dry_run": dry_run,
        "sap_write": 0,
        "excel_write": 0,
        "backup_path": None,
        "analise": analise,
        "mensagem": "Análise de reconciliação em Dry-Run concluída com sucesso (Zero escrita)."
    }

    if dry_run:
        return resultado

    # Modo Real (dry_run=False): Backup obrigatório antes de escrever
    sucesso_bkp, msg_bkp = gerar_backup_timestampado(caminho_excel)
    if not sucesso_bkp:
        resultado["ok"] = False
        resultado["mensagem"] = f"ABORTADO: Falha crítica ao gerar backup obrigatório antes da escrita: {msg_bkp}"
        return resultado

    resultado["backup_path"] = msg_bkp

    # Se não houver novas linhas para adicionar, não necessita abrir para escrita
    if analise["create_novas_count"] == 0 and analise["composta_novas_count"] == 0:
        resultado["mensagem"] = f"Nenhuma nova linha necessária. Ficheiro preservado intacto. Backup criado em: {msg_bkp}"
        return resultado

    # Gravação no Excel via openpyxl
    import openpyxl

    try:
        wb = openpyxl.load_workbook(caminho_excel)

        # 1. Gravar novas linhas em PFCG_CREATE
        if analise["create_novas_count"] > 0 and "PFCG_CREATE" in wb.sheetnames:
            ws_cr = wb["PFCG_CREATE"]
            header_cr = [_norm_col(c.value) for c in ws_cr[1]]
            col_id = next((i for i, c in enumerate(header_cr, start=1) if c == "ID"), 1)
            col_agr = next((i for i, c in enumerate(header_cr, start=1) if c in ("AGR_NAME", "ROLE")), 2)
            col_txt = next((i for i, c in enumerate(header_cr, start=1) if c in ("TEXT", "DESCRICAO")), 3)
            col_tc = next((i for i, c in enumerate(header_cr, start=1) if c in ("TCODE", "TRANSACTION")), 4)

            # Obter maior ID numérico
            max_id = 0
            for r in range(2, ws_cr.max_row + 1):
                val_id = ws_cr.cell(row=r, column=col_id).value
                if val_id is not None and str(val_id).isdigit():
                    max_id = max(max_id, int(val_id))

            for item in analise["create_novas_necessarias"]:
                max_id += 1
                row_vals = [None] * len(header_cr)
                row_vals[col_id - 1] = max_id
                row_vals[col_agr - 1] = item["role"]
                row_vals[col_txt - 1] = item["text"]
                row_vals[col_tc - 1] = item["tcode"]  # TCODE vazio
                ws_cr.append(row_vals)

        # 2. Gravar novas linhas em PFCG_COMPOSTA
        if analise["composta_novas_count"] > 0 and "PFCG_COMPOSTA" in wb.sheetnames:
            ws_cp = wb["PFCG_COMPOSTA"]
            header_cp = [_norm_col(c.value) for c in ws_cp[1]]
            col_id = next((i for i, c in enumerate(header_cp, start=1) if c == "ID"), 1)
            col_parent = next((i for i, c in enumerate(header_cp, start=1) if c in ("AGR_NAME_COMPOSTA", "COMPOSITEROLE")), 2)
            col_txt_p = next((i for i, c in enumerate(header_cp, start=1) if c in ("TEXT", "DESCRICAO")), 3)
            col_child = next((i for i, c in enumerate(header_cp, start=1) if c in ("AGR_NAME", "SINGLEROLE")), 4)

            max_id_cp = 0
            for r in range(2, ws_cp.max_row + 1):
                val_id = ws_cp.cell(row=r, column=col_id).value
                if val_id is not None and str(val_id).isdigit():
                    max_id_cp = max(max_id_cp, int(val_id))

            for item in analise["composta_novas_necessarias"]:
                max_id_cp += 1
                row_vals = [None] * len(header_cp)
                row_vals[col_id - 1] = max_id_cp
                row_vals[col_parent - 1] = item["parent"]
                row_vals[col_txt_p - 1] = item["text_parent"]
                row_vals[col_child - 1] = item["child"]
                ws_cp.append(row_vals)

        wb.save(caminho_excel)
        wb.close()

        resultado["excel_write"] = analise["create_novas_count"] + analise["composta_novas_count"]
        resultado["mensagem"] = (
            f"Reconciliação aplicada com sucesso! {analise['create_novas_count']} novas Singles em PFCG_CREATE, "
            f"{analise['composta_novas_count']} novas relações em PFCG_COMPOSTA. Backup salvo em {msg_bkp}."
        )
        return resultado

    except Exception as exc:
        resultado["ok"] = False
        resultado["mensagem"] = f"Erro ao gravar alterações no Excel: {exc}"
        return resultado
