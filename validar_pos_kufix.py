"""Validador de Pedidos de Compras (POs) - SAP PRD
Abre um popup em primeiro plano para seleção do ficheiro Excel,
extrai as POs da 1ª coluna e valida via RFC na tabela EKKO:
- Quantidade total no ficheiro
- Quantidade no SAP desmarcada (KUFIX = ' ')
- Quantidade no SAP já marcada (KUFIX = 'X')
- Eventuais POs não encontradas ou eliminadas
"""

import io
import os
import re
import sys
import msvcrt
import ctypes
from ctypes import wintypes
from pathlib import Path
from datetime import datetime

import openpyxl
from dotenv import load_dotenv

# Configuração Tkinter para interface gráfica
import tkinter as tk
from tkinter import ttk, filedialog, messagebox


# ----------------------------------------------------------------------
# 1. Leitura de Ficheiro Excel com suporte a ficheiros abertos
# ----------------------------------------------------------------------
def read_excel_bytes_shared(filepath: str) -> bytes:
    """Lê os bytes de um ficheiro mesmo que esteja aberto no Excel (Windows API)."""
    GENERIC_READ = 0x80000000
    FILE_SHARE_READ = 0x00000001
    FILE_SHARE_WRITE = 0x00000002
    FILE_SHARE_DELETE = 0x00000004
    OPEN_EXISTING = 3
    FILE_ATTRIBUTE_NORMAL = 0x80

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    handle = kernel32.CreateFileW(
        filepath,
        GENERIC_READ,
        FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE,
        None,
        OPEN_EXISTING,
        FILE_ATTRIBUTE_NORMAL,
        None,
    )
    if handle == -1 or handle == 0xFFFFFFFFFFFFFFFF:
        # Tenta leitura normal de fallback
        with open(filepath, "rb") as f:
            return f.read()

    fd = msvcrt.open_osfhandle(handle, os.O_RDONLY)
    with open(fd, "rb", closefd=True) as fp:
        return fp.read()


def extract_pos_from_excel(filepath: str) -> list[dict]:
    """Extrai as POs da primeira coluna do ficheiro Excel."""
    file_bytes = read_excel_bytes_shared(filepath)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
    sheet = wb.active

    extracted = []
    seen = set()

    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not row or not row[0]:
            continue
        val = str(row[0]).strip()
        # Procura por padrão de número de pedido (ex.: 10 dígitos ou numérico)
        digits = re.sub(r"\D", "", val)
        if len(digits) == 10 and digits.startswith(("4", "5", "6", "7", "8", "9", "0")):
            po_num = digits
            is_dup = po_num in seen
            seen.add(po_num)
            extracted.append({
                "row": row_idx,
                "ebeln": po_num,
                "is_duplicate": is_dup,
                "raw_val": val,
            })
        elif val.isdigit() and len(val) >= 5:
            po_num = val.zfill(10)
            is_dup = po_num in seen
            seen.add(po_num)
            extracted.append({
                "row": row_idx,
                "ebeln": po_num,
                "is_duplicate": is_dup,
                "raw_val": val,
            })

    wb.close()
    return extracted


# ----------------------------------------------------------------------
# 2. Conexão RFC e Validação no SAP PRD
# ----------------------------------------------------------------------
def connect_sap_prd():
    from pyrfc import Connection

    env_path = Path("C:/workspace/SapScript/.env")
    if not env_path.exists():
        env_path = Path(__file__).resolve().parent / ".env"
    load_dotenv(env_path)

    params = {
        "user": os.environ["SAP_PRD_USER"],
        "passwd": os.environ["SAP_PRD_PASSWD"],
        "ashost": os.environ["SAP_PRD_ASHOST"],
        "sysnr": os.environ["SAP_PRD_SYSNR"],
        "client": os.environ["SAP_PRD_CLIENT"],
        "lang": os.getenv("SAP_PRD_LANG", "PT").strip() or "PT",
    }
    return Connection(**params)


def validate_pos_in_sap(pos_list: list[str], progress_callback=None) -> dict[str, dict]:
    """Consulta o estado dos pedidos na tabela EKKO em lotes de 50."""
    conn = connect_sap_prd()
    unique_pos = sorted(set(pos_list))
    chunk_size = 50
    total = len(unique_pos)
    results = {}

    for i in range(0, total, chunk_size):
        chunk = unique_pos[i : i + chunk_size]
        options = []
        for idx, po in enumerate(chunk):
            prefix = "OR " if idx > 0 else ""
            options.append({"TEXT": f"{prefix}EBELN = '{po}'"})

        res = conn.call(
            "RFC_READ_TABLE",
            QUERY_TABLE="EKKO",
            DELIMITER="|",
            FIELDS=[{"FIELDNAME": f} for f in ["EBELN", "BUKRS", "WAERS", "WKURS", "KUFIX", "LOEKZ", "BEDAT"]],
            OPTIONS=options,
            ROWCOUNT=0,
        )

        for row in res.get("DATA", []):
            parts = [p.strip() for p in row.get("WA", "").split("|")]
            if len(parts) >= 5:
                ebeln = parts[0]
                results[ebeln] = {
                    "found": True,
                    "bukrs": parts[1],
                    "waers": parts[2],
                    "wkurs": parts[3],
                    "kufix": parts[4],
                    "loekz": parts[5] if len(parts) > 5 else "",
                    "bedat": parts[6] if len(parts) > 6 else "",
                }

        if progress_callback:
            progress_callback(min(i + chunk_size, total), total)

    conn.close()

    # Preencher os que não foram encontrados
    for po in unique_pos:
        if po not in results:
            results[po] = {
                "found": False,
                "bukrs": "",
                "waers": "",
                "wkurs": "",
                "kufix": "",
                "loekz": "",
                "bedat": "",
            }

    return results


# ----------------------------------------------------------------------
# 3. Interface Gráfica (Tkinter)
# ----------------------------------------------------------------------
class POValidatorApp:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Validador de Taxa Fixa (KUFIX) - SAP PRD")
        self.root.geometry("960x680")
        self.root.minsize(800, 500)
        self.root.attributes("-topmost", True)

        self.style = ttk.Style(self.root)
        self.style.theme_use("clam")

        self.file_path = None
        self.pos_data = []
        self.sap_data = {}

        self._build_ui()
        self.root.after(200, self.prompt_file_selection)

    def _build_ui(self):
        # Header frame
        header_frame = tk.Frame(self.root, bg="#1E3A8A", height=70)
        header_frame.pack(fill="x")

        title_lbl = tk.Label(
            header_frame,
            text="Validação de POs - Taxa Protegida (KUFIX) em SAP PRD",
            font=("Segoe UI", 14, "bold"),
            bg="#1E3A8A",
            fg="white",
        )
        title_lbl.pack(side="left", padx=20, pady=15)

        btn_select = tk.Button(
            header_frame,
            text="📂 Selecionar Outro Ficheiro",
            font=("Segoe UI", 10, "bold"),
            bg="#3B82F6",
            fg="white",
            relief="flat",
            padx=12,
            pady=6,
            command=self.prompt_file_selection,
        )
        btn_select.pack(side="right", padx=20, pady=15)

        # File info banner
        self.file_info_var = tk.StringVar(value="Nenhum ficheiro selecionado.")
        file_banner = tk.Label(
            self.root,
            textvariable=self.file_info_var,
            font=("Segoe UI", 9),
            bg="#F3F4F6",
            fg="#4B5563",
            anchor="w",
            padx=20,
            pady=6,
        )
        file_banner.pack(fill="x")

        # Summary Cards Frame
        cards_frame = tk.Frame(self.root, bg="#F9FAFB", pady=10)
        cards_frame.pack(fill="x", padx=15)

        self.card_total = self._create_card(cards_frame, "Total no Ficheiro", "0", "#2563EB", 0)
        self.card_unmarked = self._create_card(cards_frame, "SAP: Desmarcadas (A Marcar)", "0", "#D97706", 1)
        self.card_marked = self._create_card(cards_frame, "SAP: Já Marcadas ('X')", "0", "#059669", 2)
        self.card_errors = self._create_card(cards_frame, "Não Encontradas / Inválidas", "0", "#DC2626", 3)

        # Progress bar
        self.progress_var = tk.DoubleVar(value=0)
        self.progress_bar = ttk.Progressbar(self.root, variable=self.progress_var, maximum=100)
        self.progress_bar.pack(fill="x", padx=20, pady=4)
        self.status_lbl = tk.Label(self.root, text="Pronto para validação.", font=("Segoe UI", 9, "italic"), fg="#6B7280")
        self.status_lbl.pack(anchor="w", padx=20)

        # Treeview (Tabela detalhada)
        tree_frame = tk.Frame(self.root)
        tree_frame.pack(fill="both", expand=True, padx=20, pady=10)

        columns = ("pos", "ebeln", "bukrs", "waers", "wkurs", "kufix", "status", "action")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=12)

        self.tree.heading("pos", text="# Linha")
        self.tree.heading("ebeln", text="Nº Pedido (PO)")
        self.tree.heading("bukrs", text="Empresa")
        self.tree.heading("waers", text="Moeda")
        self.tree.heading("wkurs", text="Taxa Câmbio")
        self.tree.heading("kufix", text="Flag KUFIX")
        self.tree.heading("status", text="Estado no SAP")
        self.tree.heading("action", text="Ação Necessária")

        self.tree.column("pos", width=60, anchor="center")
        self.tree.column("ebeln", width=120, anchor="center")
        self.tree.column("bukrs", width=70, anchor="center")
        self.tree.column("waers", width=70, anchor="center")
        self.tree.column("wkurs", width=110, anchor="center")
        self.tree.column("kufix", width=90, anchor="center")
        self.tree.column("status", width=180, anchor="w")
        self.tree.column("action", width=140, anchor="center")

        # Scrollbars
        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tree_frame.grid_rowconfigure(0, weight=1)
        tree_frame.grid_columnconfigure(0, weight=1)

        # Bottom Bar
        bottom_frame = tk.Frame(self.root, pady=8)
        bottom_frame.pack(fill="x", padx=20)

        btn_export = tk.Button(
            bottom_frame,
            text="📥 Exportar Relatório Excel",
            font=("Segoe UI", 9, "bold"),
            bg="#10B981",
            fg="white",
            relief="flat",
            padx=12,
            pady=6,
            command=self.export_report,
        )
        btn_export.pack(side="left")

        btn_close = tk.Button(
            bottom_frame,
            text="Fechar",
            font=("Segoe UI", 9),
            bg="#9CA3AF",
            fg="white",
            relief="flat",
            padx=12,
            pady=6,
            command=self.root.destroy,
        )
        btn_close.pack(side="right")

    def _create_card(self, parent, title: str, initial_val: str, color: str, col: int) -> tk.Label:
        frame = tk.Frame(parent, bg="white", highlightbackground="#E5E7EB", highlightthickness=1, padx=14, pady=10)
        frame.grid(row=0, column=col, padx=6, sticky="nsew")
        parent.grid_columnconfigure(col, weight=1)

        t_lbl = tk.Label(frame, text=title, font=("Segoe UI", 9, "bold"), fg="#4B5563", bg="white")
        t_lbl.pack(anchor="w")

        val_lbl = tk.Label(frame, text=initial_val, font=("Segoe UI", 18, "bold"), fg=color, bg="white")
        val_lbl.pack(anchor="w", pady=(4, 0))
        return val_lbl

    def prompt_file_selection(self):
        # Abre janela de ficheiro em primeiro plano
        self.root.lift()
        self.root.focus_force()

        desktop = Path.home() / "Desktop"
        # Tenta OneDrive Desktop se existir
        onedrive_desktop = Path.home() / "OneDrive - Salsajeans" / "Desktop"
        initial_dir = onedrive_desktop if onedrive_desktop.exists() else desktop

        path = filedialog.askopenfilename(
            parent=self.root,
            title="Selecione o ficheiro Excel com as POs na 1ª Coluna",
            initialdir=initial_dir,
            filetypes=[("Ficheiros Excel", "*.xlsx;*.xls"), ("Todos os ficheiros", "*.*")],
        )

        if not path:
            return

        self.file_path = path
        self.file_info_var.set(f"Ficheiro: {path}")
        self.process_file_and_validate()

    def process_file_and_validate(self):
        self.status_lbl.config(text="A extrair pedidos da 1ª coluna do ficheiro Excel...")
        self.root.update()

        try:
            self.pos_data = extract_pos_from_excel(self.file_path)
        except Exception as exc:
            messagebox.showerror("Erro ao ler ficheiro", f"Não foi possível ler o ficheiro Excel:\n{exc}")
            return

        total_rows = len(self.pos_data)
        unique_pos = list({x["ebeln"] for x in self.pos_data})
        self.card_total.config(text=f"{len(unique_pos)} ({total_rows} lin)")

        if not unique_pos:
            messagebox.showwarning("Nenhum pedido encontrado", "Não foi encontrada nenhuma PO válida na primeira coluna do ficheiro.")
            return

        self.status_lbl.config(text=f"A validar {len(unique_pos)} pedidos via RFC em SAP PRD...")
        self.root.update()

        def update_progress(current, total):
            pct = (current / total) * 100
            self.progress_var.set(pct)
            self.status_lbl.config(text=f"A consultar SAP PRD: {current}/{total} pedidos ({pct:.0f}%)...")
            self.root.update()

        try:
            self.sap_data = validate_pos_in_sap(unique_pos, progress_callback=update_progress)
        except Exception as exc:
            messagebox.showerror("Erro RFC SAP", f"Falha na comunicação com o SAP PRD:\n{exc}")
            return

        self._populate_results()

    def _populate_results(self):
        # Limpar tabela
        for item in self.tree.get_children():
            self.tree.delete(item)

        unmarked_count = 0
        marked_count = 0
        error_count = 0

        # Tags para colorir linhas
        self.tree.tag_configure("unmarked", background="#FEF3C7")  # Amarelo/Laranja claro
        self.tree.tag_configure("marked", background="#D1FAE5")    # Verde claro
        self.tree.tag_configure("error", background="#FEE2E2")     # Vermelho claro

        for item in self.pos_data:
            ebeln = item["ebeln"]
            sap_info = self.sap_data.get(ebeln, {})

            if not sap_info.get("found"):
                status = "Não encontrado em PRD"
                action = "Verificar / Erro"
                tag = "error"
                error_count += 1
                kufix_val = "-"
                bukrs = "-"
                waers = "-"
                wkurs = "-"
            elif sap_info.get("loekz"):
                status = f"Eliminado no SAP ({sap_info.get('loekz')})"
                action = "Ignorar (Eliminado)"
                tag = "error"
                error_count += 1
                kufix_val = sap_info.get("kufix") or " "
                bukrs = sap_info.get("bukrs")
                waers = sap_info.get("waers")
                wkurs = sap_info.get("wkurs")
            elif sap_info.get("kufix") == "X":
                status = "Já Marcada no SAP (Fixada)"
                action = "Ignorar (Já OK)"
                tag = "marked"
                marked_count += 1
                kufix_val = "X"
                bukrs = sap_info.get("bukrs")
                waers = sap_info.get("waers")
                wkurs = sap_info.get("wkurs")
            else:
                status = "Desmarcada no SAP"
                action = "A Marcar"
                tag = "unmarked"
                unmarked_count += 1
                kufix_val = " "
                bukrs = sap_info.get("bukrs")
                waers = sap_info.get("waers")
                wkurs = sap_info.get("wkurs")

            self.tree.insert(
                "",
                "end",
                values=(
                    item["row"],
                    ebeln,
                    bukrs,
                    waers,
                    wkurs,
                    f"'{kufix_val}'",
                    status,
                    action,
                ),
                tags=(tag,),
            )

        # Atualizar Cards
        self.card_unmarked.config(text=str(unmarked_count))
        self.card_marked.config(text=str(marked_count))
        self.card_errors.config(text=str(error_count))

        self.status_lbl.config(
            text=f"Validação concluída: {unmarked_count} desmarcadas (a marcar), {marked_count} já marcadas, {error_count} com aviso/erro."
        )
        self.progress_var.set(100)

    def export_report(self):
        if not self.tree.get_children():
            messagebox.showwarning("Sem dados", "Não existem dados para exportar.")
            return

        desktop = Path.home() / "Desktop"
        onedrive_desktop = Path.home() / "OneDrive - Salsajeans" / "Desktop"
        target_dir = onedrive_desktop if onedrive_desktop.exists() else desktop

        default_filename = f"Relatorio_Validacao_KUFIX_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        save_path = filedialog.asksaveasfilename(
            parent=self.root,
            title="Guardar Relatório de Validação",
            initialdir=target_dir,
            initialfile=default_filename,
            filetypes=[("Excel", "*.xlsx")],
        )

        if not save_path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Validacao_KUFIX"

        headers = [
            "Linha Excel",
            "Nº Pedido (PO)",
            "Empresa",
            "Moeda",
            "Taxa Câmbio (WKURS)",
            "Flag KUFIX",
            "Estado no SAP",
            "Ação Necessária",
        ]
        ws.append(headers)

        for child in self.tree.get_children():
            vals = self.tree.item(child)["values"]
            ws.append(vals)

        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(save_path)
        messagebox.showinfo("Sucesso", f"Relatório gravado com sucesso em:\n{save_path}")

    def run(self):
        self.root.mainloop()


if __name__ == "__main__":
    app = POValidatorApp()
    app.run()
