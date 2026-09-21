from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import argparse
import re
import shutil
import sys
from typing import Any
import unicodedata

from openpyxl import load_workbook
from pyrfc import Connection

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from sap_rfc._rfc_common import (  # noqa: E402
    build_connection_params_for,
    find_project_root,
    load_project_env,
    make_option_eq,
    make_read_only_guard,
    make_write_guard,
    read_table,
)


ROLE_RE = re.compile(r"\b(?:Z|Y)[A-Z0-9_/\-.&]{2,}\b", re.IGNORECASE)
EXCEL_PATH = PROJECT_ROOT / "S4H_Perfis de autorização_v1.xlsx"
SYSTEMS = {
    "PRD": {"env": "PRD", "subsystem": "S4PCLNT100", "ok_col": 8},
    "QAS": {"env": "QAD", "subsystem": "S4QCLNT100", "ok_col": 9},
}
WRITE_ALLOWED_FUNCTIONS = (
    "RFC_PING",
    "RFC_READ_TABLE",
    "BAPI_USER_ACTGROUPS_ASSIGN",
    "BAPI_TRANSACTION_COMMIT",
    "BAPI_TRANSACTION_ROLLBACK",
)
WRITE_ALLOWED_TABLES = ("USR02", "AGR_USERS")


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "<na>"} else text


def norm(value: Any) -> str:
    text = clean(value).upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip()


def norm_sheet(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", norm(value))


def roles_from_cell(value: Any) -> list[str]:
    out: list[str] = []
    for match in ROLE_RE.finditer(clean(value)):
        role = norm(match.group(0))
        if role and role not in out:
            out.append(role)
    return out


def find_sheet(wb: Any, target: str) -> str:
    wanted = norm_sheet(target)
    for sheet in wb.sheetnames:
        if norm_sheet(sheet) == wanted:
            return sheet
    raise ValueError(f"Sheet '{target}' não encontrada.")


def load_definitions(wb: Any) -> dict[str, set[str]]:
    ws = wb[find_sheet(wb, "DEFINIÇÕES")]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(v) for v in rows[0]]
    dep_idx = next(i for i, h in enumerate(headers) if norm_sheet(h) == "DEPARTAMENTO")
    definitions: dict[str, set[str]] = defaultdict(set)
    for row in rows[1:]:
        dep = norm(row[dep_idx] if dep_idx < len(row) else "")
        if not dep:
            continue
        for idx, value in enumerate(row):
            if idx == dep_idx:
                continue
            definitions[dep].update(roles_from_cell(value))
    return dict(definitions)


def load_users(wb: Any, definitions: dict[str, set[str]]) -> list[dict[str, Any]]:
    ws = wb[find_sheet(wb, "Proposta Ativa")]
    users: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        user = norm(row[0] if len(row) > 0 else "")
        if not user or user in seen:
            continue
        dep = norm(row[7] if len(row) > 7 else "")
        if not dep or dep not in definitions:
            continue
        seen.add(user)
        users.append(
            {
                "user": user,
                "name": clean(row[2] if len(row) > 2 else ""),
                "department": dep,
                "roles": sorted(definitions[dep]),
            }
        )
    return users


def fetch_current_roles(connection: Any, guard: Any, username: str) -> tuple[bool, dict[str, dict[str, str]]]:
    user_rows = read_table(
        connection,
        guard,
        table_name="USR02",
        fields=["BNAME"],
        options=make_option_eq("BNAME", username),
        rowcount=1,
    )
    if not user_rows:
        return False, {}
    role_rows = read_table(
        connection,
        guard,
        table_name="AGR_USERS",
        fields=["AGR_NAME", "FROM_DAT", "TO_DAT"],
        options=make_option_eq("UNAME", username),
        rowcount=0,
    )
    current: dict[str, dict[str, str]] = {}
    for agr_name, from_dat, to_dat in role_rows:
        role = norm(agr_name)
        if not role:
            continue
        current.setdefault(
            role,
            {
                "AGR_NAME": role,
                "FROM_DAT": clean(from_dat) or datetime.now().strftime("%Y%m%d"),
                "TO_DAT": clean(to_dat) or "99991231",
            },
        )
    return True, current


def assign_missing_roles(connection: Any, guard: Any, username: str, current: dict[str, dict[str, str]], missing: list[str]) -> list[str]:
    today = datetime.now().strftime("%Y%m%d")
    payload = list(current.values())
    for role in missing:
        payload.append({"AGR_NAME": role, "FROM_DAT": today, "TO_DAT": "99991231"})
    guard.assert_function_allowed("BAPI_USER_ACTGROUPS_ASSIGN")
    result = connection.call("BAPI_USER_ACTGROUPS_ASSIGN", USERNAME=username, ACTIVITYGROUPS=payload)
    errors = [
        clean(row.get("MESSAGE"))
        for row in (result.get("RETURN") or [])
        if norm(row.get("TYPE")) in {"E", "A"} and clean(row.get("MESSAGE"))
    ]
    if errors:
        try:
            connection.call("BAPI_TRANSACTION_ROLLBACK")
        except Exception:
            pass
        return errors
    guard.assert_function_allowed("BAPI_TRANSACTION_COMMIT")
    connection.call("BAPI_TRANSACTION_COMMIT", WAIT="X")
    return []


def append_cua_rows(excel_path: Path, assignments: list[dict[str, str]]) -> None:
    if not assignments:
        return
    wb = load_workbook(excel_path)
    ws = wb[find_sheet(wb, "CUA_ADICIONAR")]
    max_id = 0
    for row in range(2, ws.max_row + 1):
        try:
            max_id = max(max_id, int(ws.cell(row, 1).value or 0))
        except (TypeError, ValueError):
            pass
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for item in assignments:
        max_id += 1
        new_row = ws.max_row + 1
        ws.cell(new_row, 1, max_id)
        ws.cell(new_row, 2, item["user"])
        ws.cell(new_row, 3, item["subsystem"])
        ws.cell(new_row, 4, item["role"])
        ws.cell(new_row, 5, "CONCLUÍDO")
        ws.cell(new_row, 6, f"Atribuição acrescentada por regra DEFINIÇÕES em {item['label']} (add-only).")
        ws.cell(new_row, 7, ts)
        ws.cell(new_row, int(item["ok_col"]), "OK")
    wb.save(excel_path)
    wb.close()


def run(execute: bool) -> int:
    load_project_env(find_project_root())
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    definitions = load_definitions(wb)
    users = load_users(wb, definitions)
    wb.close()

    print(f"EXCEL={EXCEL_PATH}")
    print(f"MODE={'EXECUTE' if execute else 'PLAN'}")
    print(f"DEPARTMENTS={len(definitions)}")
    print(f"USERS_IN_SCOPE={len(users)}")

    if execute:
        backup_dir = PROJECT_ROOT / "output"
        backup_dir.mkdir(exist_ok=True)
        backup_path = backup_dir / f"S4H_Perfis_backup_pre_definicoes_add_only_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        shutil.copy2(EXCEL_PATH, backup_path)
        print(f"BACKUP={backup_path}")

    successful_assignments: list[dict[str, str]] = []
    for label, system in SYSTEMS.items():
        params = build_connection_params_for(system["env"])
        guard = make_write_guard(WRITE_ALLOWED_FUNCTIONS, WRITE_ALLOWED_TABLES)
        read_guard = make_read_only_guard(WRITE_ALLOWED_TABLES)
        print(f"SYSTEM={label} SUBSYSTEM={system['subsystem']}")
        connection = Connection(**params)
        try:
            connection.call("RFC_PING")
            system_missing = 0
            system_users = 0
            for item in users:
                exists, current = fetch_current_roles(connection, read_guard, item["user"])
                if not exists:
                    print(f"  SKIP_NOT_FOUND {item['user']} {item['department']}")
                    continue
                missing = [role for role in item["roles"] if role not in current]
                if not missing:
                    continue
                system_users += 1
                system_missing += len(missing)
                print(f"  MISSING {item['user']} {item['department']} {len(missing)}: {', '.join(missing)}")
                if execute:
                    errors = assign_missing_roles(connection, guard, item["user"], current, missing)
                    if errors:
                        print(f"  ERROR {item['user']}: {'; '.join(errors)}")
                        continue
                    for role in missing:
                        successful_assignments.append(
                            {
                                "user": item["user"],
                                "subsystem": system["subsystem"],
                                "role": role,
                                "label": label,
                                "ok_col": str(system["ok_col"]),
                            }
                        )
                    print(f"  ADDED {item['user']} {len(missing)}")
            print(f"SYSTEM_SUMMARY {label}: USERS_WITH_MISSING={system_users} MISSING_ASSIGNMENTS={system_missing}")
        finally:
            try:
                connection.close()
            except Exception:
                pass

    if execute:
        append_cua_rows(EXCEL_PATH, successful_assignments)
        print(f"CUA_ADICIONAR_ROWS_APPENDED={len(successful_assignments)}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Aplica add-only das roles da sheet DEFINIÇÕES por departamento.")
    parser.add_argument("--execute", action="store_true", help="Executa BAPI e acrescenta linhas em CUA_ADICIONAR.")
    args = parser.parse_args()
    return run(args.execute)


if __name__ == "__main__":
    raise SystemExit(main())
