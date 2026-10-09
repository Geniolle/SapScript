from pathlib import Path
import time

import openpyxl
import win32com.client


# ============================================================
# CONFIGURAÇÃO
# ============================================================

ARQUIVO_EXCEL = Path(r"C:\Temp\OB09_MASSA.xlsx")

PLANO_CONTAS = "1000"

# True  = apenas valida/compara, não grava
# False = altera e grava na OB09
MODO_SIMULACAO = True

TEMPO_ESPERA = 0.3


# ============================================================
# COLUNAS DO EXCEL
# ============================================================

COL_CONTA = 1
COL_LSBEW = 2
COL_LHBEW = 3
COL_LHTRV = 4
COL_LSTRV = 5

COL_STATUS = 6
COL_MENSAGEM = 7

COL_ATUAL_LSBEW = 8
COL_ATUAL_LHBEW = 9
COL_ATUAL_LHTRV = 10
COL_ATUAL_LSTRV = 11


# ============================================================
# IDS SAP
# ============================================================

ID_POSICIONAR = "wnd[0]/usr/btnVIM_POSI_PUSH"

ID_CONTA_TABELA = (
    "wnd[0]/usr/tblSAPL0F11TCTRL_V_T030H/"
    "ctxtV_T030H-HKONT[0,0]"
)

ID_LSBEW = "wnd[0]/usr/ctxtV_T030H-LSBEW"
ID_LHBEW = "wnd[0]/usr/ctxtV_T030H-LHBEW"
ID_LHTRV = "wnd[0]/usr/ctxtV_T030H-LHTRV"
ID_LSTRV = "wnd[0]/usr/ctxtV_T030H-LSTRV"


# ============================================================
# SAP
# ============================================================

def conectar_sap():
    """
    Liga-se a uma sessão SAP GUI já aberta.
    """

    sap_gui_auto = win32com.client.GetObject("SAPGUI")
    application = sap_gui_auto.GetScriptingEngine

    if application.Children.Count == 0:
        raise RuntimeError("Nenhuma ligação SAP aberta.")

    connection = application.Children(0)

    if connection.Children.Count == 0:
        raise RuntimeError("Nenhuma sessão SAP encontrada.")

    session = connection.Children(0)

    return session


def esperar():
    time.sleep(TEMPO_ESPERA)


def abrir_ob09(session):
    print("Abrindo OB09...")

    session.findById("wnd[0]/tbar[0]/okcd").text = "/nOB09"
    session.findById("wnd[0]").sendVKey(0)

    esperar()

    # Popup inicial
    campo_plano = (
        "wnd[1]/usr/sub:SAPLSVIX:0100/"
        "ctxtD0100_FIELD_TAB-LOWER_LIMIT[0,37]"
    )

    session.findById(campo_plano).text = PLANO_CONTAS
    session.findById("wnd[1]").sendVKey(0)

    esperar()


def localizar_conta(session, conta):
    """
    Usa o botão Posicionar da OB09.
    """

    session.findById(ID_POSICIONAR).press()
    esperar()

    campo_busca = (
        "wnd[1]/usr/sub:SAPLSPO4:0300/"
        "ctxtSVALD-VALUE[0,21]"
    )

    session.findById(campo_busca).text = conta
    session.findById("wnd[1]/tbar[0]/btn[0]").press()

    esperar()

    # Validação de segurança
    conta_encontrada = (
        session.findById(ID_CONTA_TABELA)
        .text
        .strip()
    )

    if conta_encontrada != conta:
        raise RuntimeError(
            f"Conta encontrada '{conta_encontrada}' "
            f"é diferente da solicitada '{conta}'."
        )

    return conta_encontrada


def abrir_detalhe(session):
    session.findById(ID_CONTA_TABELA).setFocus()
    session.findById("wnd[0]").sendVKey(2)

    esperar()


def voltar_lista(session):
    session.findById("wnd[0]/tbar[0]/btn[3]").press()
    esperar()


def ler_valores_atuais(session):
    return {
        "LSBEW": session.findById(ID_LSBEW).text.strip(),
        "LHBEW": session.findById(ID_LHBEW).text.strip(),
        "LHTRV": session.findById(ID_LHTRV).text.strip(),
        "LSTRV": session.findById(ID_LSTRV).text.strip(),
    }


def alterar_valores(
    session,
    lsbew,
    lhbew,
    lhtrv,
    lstrv,
):
    session.findById(ID_LSBEW).text = lsbew
    session.findById(ID_LHBEW).text = lhbew
    session.findById(ID_LHTRV).text = lhtrv
    session.findById(ID_LSTRV).text = lstrv


def gravar(session):
    """
    Ctrl+S / botão Save.
    """

    session.findById("wnd[0]/tbar[0]/btn[11]").press()
    esperar()


# ============================================================
# EXCEL
# ============================================================

def valor_excel(valor):
    """
    Evita valores como 11300107.0.
    """

    if valor is None:
        return ""

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor).strip()


def preparar_excel():
    if not ARQUIVO_EXCEL.exists():
        raise FileNotFoundError(
            f"Excel não encontrado:\n{ARQUIVO_EXCEL}"
        )

    workbook = openpyxl.load_workbook(ARQUIVO_EXCEL)
    sheet = workbook.active

    # Cabeçalhos de controlo
    sheet.cell(1, COL_STATUS).value = "STATUS"
    sheet.cell(1, COL_MENSAGEM).value = "MENSAGEM"

    sheet.cell(1, COL_ATUAL_LSBEW).value = "ATUAL_LSBEW"
    sheet.cell(1, COL_ATUAL_LHBEW).value = "ATUAL_LHBEW"
    sheet.cell(1, COL_ATUAL_LHTRV).value = "ATUAL_LHTRV"
    sheet.cell(1, COL_ATUAL_LSTRV).value = "ATUAL_LSTRV"

    workbook.save(ARQUIVO_EXCEL)

    return workbook, sheet


# ============================================================
# PROCESSAMENTO
# ============================================================

def processar_linha(session, sheet, linha):
    conta = valor_excel(
        sheet.cell(linha, COL_CONTA).value
    )

    lsbew = valor_excel(
        sheet.cell(linha, COL_LSBEW).value
    )

    lhbew = valor_excel(
        sheet.cell(linha, COL_LHBEW).value
    )

    lhtrv = valor_excel(
        sheet.cell(linha, COL_LHTRV).value
    )

    lstrv = valor_excel(
        sheet.cell(linha, COL_LSTRV).value
    )

    if not conta:
        return

    print()
    print("=" * 70)
    print(f"Linha: {linha}")
    print(f"Conta: {conta}")

    localizar_conta(
        session=session,
        conta=conta,
    )

    abrir_detalhe(session)

    atuais = ler_valores_atuais(session)

    # Guardar valores existentes
    sheet.cell(
        linha,
        COL_ATUAL_LSBEW,
    ).value = atuais["LSBEW"]

    sheet.cell(
        linha,
        COL_ATUAL_LHBEW,
    ).value = atuais["LHBEW"]

    sheet.cell(
        linha,
        COL_ATUAL_LHTRV,
    ).value = atuais["LHTRV"]

    sheet.cell(
        linha,
        COL_ATUAL_LSTRV,
    ).value = atuais["LSTRV"]

    novos = {
        "LSBEW": lsbew,
        "LHBEW": lhbew,
        "LHTRV": lhtrv,
        "LSTRV": lstrv,
    }

    print("Valores atuais:")
    print(atuais)

    print("Novos valores:")
    print(novos)

    if atuais == novos:
        print("Nenhuma alteração necessária.")

        sheet.cell(
            linha,
            COL_STATUS,
        ).value = "SEM ALTERAÇÃO"

        sheet.cell(
            linha,
            COL_MENSAGEM,
        ).value = "Valores já estão corretos."

        voltar_lista(session)

        return

    if MODO_SIMULACAO:
        print("SIMULAÇÃO - nenhuma gravação realizada.")

        sheet.cell(
            linha,
            COL_STATUS,
        ).value = "SIMULAÇÃO"

        sheet.cell(
            linha,
            COL_MENSAGEM,
        ).value = "Alteração necessária, não gravada."

        voltar_lista(session)

        return

    alterar_valores(
        session=session,
        lsbew=lsbew,
        lhbew=lhbew,
        lhtrv=lhtrv,
        lstrv=lstrv,
    )

    gravar(session)

    # ========================================================
    # READBACK
    # ========================================================

    valores_depois = ler_valores_atuais(session)

    if valores_depois != novos:
        raise RuntimeError(
            "Valores após gravação são diferentes "
            "dos valores esperados."
        )

    sheet.cell(
        linha,
        COL_STATUS,
    ).value = "OK"

    sheet.cell(
        linha,
        COL_MENSAGEM,
    ).value = "Alteração gravada e validada."

    print("OK - alteração gravada.")

    voltar_lista(session)


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 70)
    print("OB09 - ALTERAÇÃO EM MASSA")
    print("=" * 70)

    if MODO_SIMULACAO:
        print("MODO: SIMULAÇÃO")
        print("Nenhuma alteração será gravada.")
    else:
        print("MODO: EXECUÇÃO REAL")
        print("As alterações serão gravadas no SAP.")

    print()

    workbook, sheet = preparar_excel()

    session = conectar_sap()

    abrir_ob09(session)

    ultima_linha = sheet.max_row

    total = 0
    sucesso = 0
    erros = 0

    for linha in range(2, ultima_linha + 1):

        conta = valor_excel(
            sheet.cell(linha, COL_CONTA).value
        )

        if not conta:
            continue

        total += 1

        try:
            processar_linha(
                session=session,
                sheet=sheet,
                linha=linha,
            )

            sucesso += 1

        except Exception as erro:
            erros += 1

            print(
                f"ERRO conta {conta}: {erro}"
            )

            sheet.cell(
                linha,
                COL_STATUS,
            ).value = "ERRO"

            sheet.cell(
                linha,
                COL_MENSAGEM,
            ).value = str(erro)

            # tentativa de recuperação
            try:
                session.findById(
                    "wnd[0]/tbar[0]/btn[3]"
                ).press()
            except Exception:
                pass

        # Guarda após cada conta
        workbook.save(ARQUIVO_EXCEL)

    workbook.save(ARQUIVO_EXCEL)

    print()
    print("=" * 70)
    print("PROCESSAMENTO CONCLUÍDO")
    print("=" * 70)

    print(f"Total:   {total}")
    print(f"OK:      {sucesso}")
    print(f"Erros:   {erros}")

    print()
    print(f"Resultado Excel: {ARQUIVO_EXCEL}")


if __name__ == "__main__":
    main()