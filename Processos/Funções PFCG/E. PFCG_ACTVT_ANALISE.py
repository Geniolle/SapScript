# -*- coding: utf-8 -*-
"""
E. PFCG_ACTVT_ANALISE.py
======================================================================
Somente leitura, via RFC: para as funções (AGR_NAME únicos) da sheet
'PFCG_CREATE' do ficheiro Excel oficial, verifica em SAP PRD quais
delas têm o campo ACTVT preenchido em algum objeto de autorização
(AGR_1251.FIELD = 'ACTVT'), expandindo funções compostas (Sammelrolle)
para as funções-membro.

Não altera o SAP. Apenas lê (RFC_READ_TABLE) em PRD.
======================================================================
"""

import sys
from pathlib import Path
from typing import Any, Dict, List

FOLDER = Path(__file__).resolve().parent
PROJECT_ROOT = FOLDER.parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

CAMINHO_EXCEL_PADRAO = PROJECT_ROOT / "sap_script_uploads" / "S4H_Perfis de autorização.xlsx"
NOME_SHEET = "PFCG_CREATE"


def obter_roles_pfcg_create(caminho_excel: str = None) -> List[str]:
    from openpyxl import load_workbook

    caminho = Path(caminho_excel) if caminho_excel else CAMINHO_EXCEL_PADRAO
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb[NOME_SHEET]
        linhas = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not linhas:
        return []

    cabecalho = linhas[0]
    idx_agr = cabecalho.index("AGR_NAME")

    roles: List[str] = []
    vistos = set()
    for linha in linhas[1:]:
        valor = linha[idx_agr]
        if valor is None:
            continue
        nome = str(valor).strip().upper()
        if nome and nome not in vistos:
            vistos.add(nome)
            roles.append(nome)
    return roles


def analisar_actvt(roles: List[str]) -> Dict[str, Any]:
    from sap_rfc._rfc_common import (
        build_connection_params_for, load_project_env, find_project_root,
        make_read_only_guard, make_option_in, read_table,
    )
    from pyrfc import Connection

    project_root = find_project_root()
    load_project_env(project_root)
    params = build_connection_params_for("PRD")
    guard = make_read_only_guard(("AGR_DEFINE", "AGR_1251", "AGR_AGRS"))
    conn = Connection(**params)
    try:
        # Expande funções compostas (Sammelrolle) para as funções-membro.
        agrs_rows = read_table(
            conn, guard,
            table_name="AGR_AGRS",
            fields=["AGR_NAME", "CHILD_AGR"],
            options=make_option_in("AGR_NAME", roles),
            rowcount=0,
        )
        membros_por_role: Dict[str, List[str]] = {}
        for agr_name, child_agr in agrs_rows:
            membros_por_role.setdefault(agr_name.strip().upper(), []).append(child_agr.strip().upper())

        todas_as_roles_a_ler = set(roles)
        for membros in membros_por_role.values():
            todas_as_roles_a_ler.update(membros)

        auth_rows = read_table(
            conn, guard,
            table_name="AGR_1251",
            fields=["AGR_NAME", "OBJECT", "FIELD", "DELETED"],
            options=make_option_in("AGR_NAME", sorted(todas_as_roles_a_ler)),
            rowcount=0,
        )
    finally:
        conn.close()

    # Para cada role (ou função-membro, se composta), lista os objetos onde ACTVT aparece.
    objetos_actvt_por_role: Dict[str, set] = {}
    for agr_name, obj, campo, deleted in auth_rows:
        if str(deleted).strip().upper() == "X":
            continue
        if campo.strip().upper() != "ACTVT":
            continue
        agr_name = agr_name.strip().upper()
        objetos_actvt_por_role.setdefault(agr_name, set()).add(obj.strip())

    resultado: List[Dict[str, Any]] = []
    for role in roles:
        membros = membros_por_role.get(role, [])
        roles_a_verificar = [role] + membros
        objetos = set()
        for r in roles_a_verificar:
            objetos.update(objetos_actvt_por_role.get(r, set()))

        resultado.append({
            "AGR_NAME": role,
            "is_composite": bool(membros),
            "composite_members": membros,
            "tem_actvt": bool(objetos),
            "objetos_com_actvt": sorted(objetos),
        })

    return {
        "system": "PRD",
        "client": params["client"],
        "total_roles": len(roles),
        "resultado": resultado,
    }


def main(caminho_excel: str = None):
    print("=" * 75)
    print("  ANALISE (RFC, SOMENTE LEITURA) — ACTVT NAS FUNÇÕES DA SHEET PFCG_CREATE")
    print("=" * 75)

    roles = obter_roles_pfcg_create(caminho_excel)
    print(f"\nFunções únicas encontradas em '{NOME_SHEET}': {len(roles)}")

    print("\nA consultar SAP PRD (AGR_1251, AGR_AGRS) via RFC...")
    analise = analisar_actvt(roles)

    com_actvt = [r for r in analise["resultado"] if r["tem_actvt"]]
    sem_actvt = [r for r in analise["resultado"] if not r["tem_actvt"]]

    print("\n" + "-" * 75)
    print(f"RESUMO (Sistema: {analise['system']} | Client: {analise['client']})")
    print("-" * 75)
    print(f"Total de funções analisadas: {analise['total_roles']}")
    print(f"Com campo ACTVT preenchido: {len(com_actvt)}")
    print(f"Sem campo ACTVT: {len(sem_actvt)}")

    print(f"\nFunções COM ACTVT ({len(com_actvt)}):")
    for r in com_actvt:
        tag = " [composta]" if r["is_composite"] else ""
        print(f"  - {r['AGR_NAME']}{tag}: objetos -> {', '.join(r['objetos_com_actvt'])}")

    print(f"\nFunções SEM ACTVT ({len(sem_actvt)}):")
    for r in sem_actvt:
        tag = " [composta]" if r["is_composite"] else ""
        print(f"  - {r['AGR_NAME']}{tag}")

    return analise


if __name__ == "__main__":
    caminho = sys.argv[1] if len(sys.argv) > 1 else None
    main(caminho)
