# -*- coding: utf-8 -*-

###################################################################################
# N. PESQUISAR_TRANSACAO.py
# PFCG - Pesquisar a função individual de uma transação SAP e o respetivo estado
# de atribuição no Excel (composites, utilizadores, sheets de departamento), com
# verificação opcional em SAP via RFC (só leitura).
#
# Fluxo:
#   1. Sheet "Proposta": localizar o bloco de função individual (Z_...) que contém
#      a transação indicada (mesma lógica de M. PFCG_ADD_FUNCAO_COMPOSTA.py).
#   2. Sheet "PFCG_COMPOSTA": listar as composites que já incluem essa função
#      individual (segundo o Excel).
#   3. Sheet "Proposta Ativa": listar os utilizadores que já têm a função
#      individual, direta ou indiretamente (composta associada).
#   4. Sheets de departamento (matrizes por equipa, detetadas dinamicamente pelo
#      cabeçalho "Transação"/"Descrição" na 2ª linha): localizar a linha da
#      transação, se existir, e listar quem tem 'X'.
#   5. (Opcional, --ambiente) RFC AGR_DEFINE + AGR_AGRS: confirmar em SAP se a
#      função individual existe e se cada composta encontrada no Excel já a tem
#      como membro. O Excel já divergiu do SAP real nesta pasta (writes perdidos
#      por conflito de sincronização do SharePoint) — usar isto para confirmar.
###################################################################################

import sys
import argparse
from pathlib import Path
from typing import Any, Dict, List, Optional, Set

FOLDER = Path(__file__).resolve().parent
PROJECT_ROOT = FOLDER.parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

CAMINHO_EXCEL_PADRAO = PROJECT_ROOT / "sap_script_uploads" / "S4H_Perfis de autorização.xlsx"
NOME_SHEET_PROPOSTA = "Proposta"
NOME_SHEET_PROPOSTA_ATIVA = "Proposta Ativa"
NOME_SHEET_PFCG_COMPOSTA = "PFCG_COMPOSTA"

ALLOWED_TABLES = ("AGR_AGRS", "AGR_DEFINE", "TSTCT")
ALLOWED_FUNCTIONS = ("RFC_READ_TABLE",)


def localizar_funcao_individual_por_tcode(tcode: str, caminho_excel: Optional[str] = None) -> Optional[str]:
    """Sheet 'Proposta': encontra a função individual (Z_...) cujo bloco contém a tcode."""
    import openpyxl

    tcode = tcode.strip().upper()
    caminho = Path(caminho_excel) if caminho_excel else CAMINHO_EXCEL_PADRAO
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb[NOME_SHEET_PROPOSTA]

    role_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=2).value]
    role_rows.append(ws.max_row + 1)

    for i in range(len(role_rows) - 1):
        inicio, fim = role_rows[i], role_rows[i + 1]
        role_name = ws.cell(row=inicio, column=1).value
        tcodes_bloco = {
            str(ws.cell(row=r, column=1).value).strip().upper()
            for r in range(inicio + 1, fim)
            if ws.cell(row=r, column=1).value
        }
        if tcode in tcodes_bloco:
            return str(role_name).strip().upper()
    return None


def listar_composites_com_funcao(funcao_individual: str, caminho_excel: Optional[str] = None) -> List[Dict[str, Any]]:
    """Sheet 'PFCG_COMPOSTA': composites que já incluem a função individual, segundo o Excel."""
    import openpyxl

    caminho = Path(caminho_excel) if caminho_excel else CAMINHO_EXCEL_PADRAO
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb[NOME_SHEET_PFCG_COMPOSTA]

    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_composta = next((i + 1 for i, h in enumerate(header) if "COMPOSTA" in str(h or "").strip().upper()), 2)
    col_texto = next((i + 1 for i, h in enumerate(header) if str(h or "").strip().upper() == "TEXT"), 3)
    col_agr_name = next((i + 1 for i, h in enumerate(header) if str(h or "").strip().upper() == "AGR_NAME"), 4)
    col_status = next((i + 1 for i, h in enumerate(header) if str(h or "").strip().upper() == "STATUS"), 5)

    resultados = []
    for r in range(2, ws.max_row + 1):
        agr_name = ws.cell(row=r, column=col_agr_name).value
        if str(agr_name or "").strip().upper() == funcao_individual:
            resultados.append({
                "composta": str(ws.cell(row=r, column=col_composta).value or "").strip(),
                "texto_composta": str(ws.cell(row=r, column=col_texto).value or "").strip(),
                "status": str(ws.cell(row=r, column=col_status).value or "").strip(),
            })
    return resultados


def listar_utilizadores_com_funcao(
    funcao_individual: str,
    composites: Set[str],
    caminho_excel: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Sheet 'Proposta Ativa': utilizadores com a função individual, direta ou via composta."""
    import openpyxl

    caminho = Path(caminho_excel) if caminho_excel else CAMINHO_EXCEL_PADRAO
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb[NOME_SHEET_PROPOSTA_ATIVA]

    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_usuario = next((i + 1 for i, h in enumerate(header) if str(h or "").strip().upper() == "USUÁRIO"), 1)
    col_nome = next((i + 1 for i, h in enumerate(header) if "NOME" in str(h or "").strip().upper()), 3)
    col_departamento = next((i + 1 for i, h in enumerate(header) if str(h or "").strip().upper() == "DEPARTAMENTO"), 8)
    col_composite = next((i + 1 for i, h in enumerate(header) if "COMPOSITE" in str(h or "").strip().upper()), 9)
    col_desc = next((i + 1 for i, h in enumerate(header) if "DESCRI" in str(h or "").strip().upper()), 10)
    inicio_roles = col_desc + 1

    resultados = []
    for r in range(2, ws.max_row + 1):
        usuario = ws.cell(row=r, column=col_usuario).value
        if not usuario:
            continue
        composta_user = str(ws.cell(row=r, column=col_composite).value or "").strip().upper()
        roles_user = {
            str(ws.cell(row=r, column=c).value or "").strip().upper()
            for c in range(inicio_roles, ws.max_column + 1)
            if ws.cell(row=r, column=c).value
        }

        tem_direto = funcao_individual in roles_user
        tem_via_composta = composta_user in composites
        if not (tem_direto or tem_via_composta):
            continue

        resultados.append({
            "usuario": str(usuario).strip(),
            "nome": str(ws.cell(row=r, column=col_nome).value or "").strip(),
            "departamento": str(ws.cell(row=r, column=col_departamento).value or "").strip(),
            "composite_role": composta_user,
            "origem": "Direto" if tem_direto else "Via composta",
        })
    return resultados


def _eh_sheet_departamento(ws) -> bool:
    """Deteta dinamicamente as sheets de matriz por departamento pelo cabeçalho da linha 2."""
    col_a = str(ws.cell(row=2, column=1).value or "").strip().upper()
    col_b = str(ws.cell(row=2, column=2).value or "").strip().upper()
    return "TRANSA" in col_a and "DESCRI" in col_b


def localizar_nas_sheets_departamento(tcode: str, caminho_excel: Optional[str] = None) -> Dict[str, Dict[str, Any]]:
    """Percorre todas as sheets de departamento (matriz) e reporta a linha da tcode, se existir."""
    import openpyxl

    tcode = tcode.strip().upper()
    caminho = Path(caminho_excel) if caminho_excel else CAMINHO_EXCEL_PADRAO
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)

    resultados: Dict[str, Dict[str, Any]] = {}
    for nome_sheet in wb.sheetnames:
        ws = wb[nome_sheet]
        if not _eh_sheet_departamento(ws):
            continue

        header_users = [ws.cell(row=2, column=c).value for c in range(3, ws.max_column + 1)]
        linha_tcode = None
        for r in range(3, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value or "").strip().upper() == tcode:
                linha_tcode = r
                break

        if linha_tcode is None:
            resultados[nome_sheet] = {"encontrado": False}
            continue

        descricao = str(ws.cell(row=linha_tcode, column=2).value or "").strip()
        utilizadores_com_x = []
        for i, h in enumerate(header_users):
            if not h:
                continue
            valor = ws.cell(row=linha_tcode, column=3 + i).value
            if str(valor or "").strip().upper() == "X":
                utilizadores_com_x.append(str(h).split("\n")[0].strip())

        resultados[nome_sheet] = {
            "encontrado": True,
            "linha": linha_tcode,
            "descricao": descricao,
            "utilizadores_com_x": utilizadores_com_x,
        }
    return resultados


def verificar_no_sap(ambiente: str, funcao_individual: str, composites: List[str]) -> Dict[str, Any]:
    """RFC (somente leitura): confirma se a função existe e quais composites já a têm como membro."""
    from sap_rfc._rfc_common import (
        build_connection_params_for, load_project_env, find_project_root,
        make_read_only_guard, role_exists, fetch_composite_members,
    )
    from pyrfc import Connection

    load_project_env(find_project_root())
    params = build_connection_params_for(ambiente)
    conn = Connection(**params)
    guard = make_read_only_guard(ALLOWED_TABLES)
    try:
        existe = role_exists(conn, guard, funcao_individual)
        membros_por_composta = {}
        for c in composites:
            membros = fetch_composite_members(conn, guard, c)
            membros_por_composta[c] = funcao_individual in membros
    finally:
        conn.close()

    return {"funcao_existe_sap": existe, "membros_por_composta": membros_por_composta}


def executar_pesquisa(tcode: str, caminho_excel: Optional[str] = None, ambiente: Optional[str] = None) -> Dict[str, Any]:
    tcode = tcode.strip().upper()

    funcao_individual = localizar_funcao_individual_por_tcode(tcode, caminho_excel)
    if not funcao_individual:
        return {"ok": False, "tcode": tcode, "message": f"Nenhuma função individual encontrada para a transação '{tcode}' na sheet '{NOME_SHEET_PROPOSTA}'."}

    composites_info = listar_composites_com_funcao(funcao_individual, caminho_excel)
    composites_nomes = {c["composta"].upper() for c in composites_info}

    utilizadores = listar_utilizadores_com_funcao(funcao_individual, composites_nomes, caminho_excel)
    sheets_departamento = localizar_nas_sheets_departamento(tcode, caminho_excel)

    resultado: Dict[str, Any] = {
        "ok": True,
        "tcode": tcode,
        "funcao_individual": funcao_individual,
        "composites": composites_info,
        "utilizadores": utilizadores,
        "sheets_departamento": sheets_departamento,
    }

    if ambiente:
        resultado["sap"] = verificar_no_sap(ambiente, funcao_individual, sorted(composites_nomes))

    return resultado


def imprimir_resultado(resultado: Dict[str, Any]) -> None:
    print("\n" + "=" * 75)
    print(f"PESQUISAR TRANSAÇÃO: {resultado['tcode']}")
    print("=" * 75)

    if not resultado["ok"]:
        print(f"[ERRO] {resultado['message']}")
        return

    print(f"Função individual: {resultado['funcao_individual']}")

    print(f"\nCOMPOSITES COM ESTA FUNÇÃO ({len(resultado['composites'])}) [Excel/PFCG_COMPOSTA]:")
    if resultado["composites"]:
        for c in resultado["composites"]:
            print(f"  - {c['composta']:<35} | {c['texto_composta']:<40} | {c['status']}")
    else:
        print("  Nenhuma composta encontrada em PFCG_COMPOSTA.")

    print(f"\nUTILIZADORES COM ACESSO ({len(resultado['utilizadores'])}) [Excel/Proposta Ativa]:")
    if resultado["utilizadores"]:
        for u in resultado["utilizadores"]:
            print(f"  - {u['usuario']:<12} | {u['nome']:<25} | {u['departamento']:<28} | {u['composite_role']:<32} | {u['origem']}")
    else:
        print("  Nenhum utilizador encontrado.")

    print("\nSHEETS DE DEPARTAMENTO (matriz por equipa):")
    for nome_sheet, info in resultado["sheets_departamento"].items():
        if not info["encontrado"]:
            print(f"  [{nome_sheet}] - transação não consta nesta sheet.")
            continue
        users = ", ".join(info["utilizadores_com_x"]) if info["utilizadores_com_x"] else "(nenhum com X)"
        print(f"  [{nome_sheet}] linha {info['linha']} | '{info['descricao']}' | Com X: {users}")

    if "sap" in resultado:
        sap = resultado["sap"]
        print(f"\nVERIFICAÇÃO EM SAP ({'existe' if sap['funcao_existe_sap'] else 'NÃO EXISTE'}):")
        for composta, membro in sap["membros_por_composta"].items():
            print(f"  - {composta:<35} | membro em SAP? {membro}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Pesquisar a função individual de uma transação SAP e o seu estado de atribuição.")
    parser.add_argument("--tcode", required=False, help="Transação (ex.: KS03). Se omitido, é pedida interativamente.")
    parser.add_argument("--excel", dest="excel", help="Caminho do ficheiro Excel (default: sap_script_uploads/S4H_Perfis de autorização.xlsx).")
    parser.add_argument("--ambiente", choices=["DEV", "QAD", "PRD"], help="Se indicado, confirma via RFC (somente leitura) se a função existe e é membro das composites encontradas.")
    args = parser.parse_args()

    tcode = args.tcode
    while not tcode:
        tcode = input("Qual a transação SAP a analisar (ex.: KS03)? ").strip()

    resultado = executar_pesquisa(tcode, args.excel, args.ambiente)
    imprimir_resultado(resultado)
