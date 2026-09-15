# -*- coding: utf-8 -*-

###################################################################################
# L. PFCG_ELIMINAR_CAMPO.py
# PFCG - Eliminar autorização global de um campo (ex.: BUKRS) via PFCGMASSVAL
#
# Diferente de "F. PFCG_ACTVT_GLOBAL.py" (que substitui o valor de um campo por
# '*' — Full Auth), este script usa o modo "D - Eliminar" para remover qualquer
# valor real que o campo tenha, deixando-o sem valor (placeholder de nível
# organizacional '$<CAMPO>', que representa "não maintido / sem valor real").
#
# Método validado (2026-09-13/14) em Simulação e confirmado pelo utilizador:
#   1. PFCGMASSVAL, Execução Direta, radSEL_FLD ("Campo para vários objetos")
#   2. ROLE-LOW = função; cmbFLDACT = 'D' (Eliminar); ctxtFLDOBJ em branco
#      (todos os objetos); ctxtFLDFLD = campo (ex.: BUKRS)
#   3. chkS_ACHD desmarcada (mesmo motivo que em F. PFCG_ACTVT_GLOBAL.py)
#   4. Botão "Vals." (btnPFLDN) -> popup "Valores de campo" -> botão
#      "Autorização global para" (btnGES1) -> confirmar "Sim" no popup
#      "Inserir autorização global" -> confirmar popup "Valores de campo"
#   5. Executar (F8) -> "Gerar perfis" se aparecer
#
# Verificação (RFC, AGR_1251): sucesso = nenhuma linha não eliminada do campo
# fica com um valor "real" (qualquer coisa diferente do placeholder
# '$<CAMPO>' já é considerado valor real e é reportado como problema).
###################################################################################

import sys
import time
from pathlib import Path
from typing import Any, Dict, List

FOLDER = Path(__file__).resolve().parent
PROJECT_ROOT = FOLDER.parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

MAPA_SISTEMA = {"DEV": "S4D", "QAD": "S4Q", "PRD": "S4P"}
CAMINHO_EXCEL_PADRAO = PROJECT_ROOT / "sap_script_uploads" / "S4H_Perfis de autorização.xlsx"
NOME_SHEET_OBJETOS_ATUALIZADOS = "OBJETOS ATUALIZADOS"


def registar_objetos_atualizados(role_name: str, objetos_valores: List[tuple], caminho_excel: str = None) -> None:
    from openpyxl import load_workbook

    if not objetos_valores:
        return

    caminho = Path(caminho_excel) if caminho_excel else CAMINHO_EXCEL_PADRAO
    wb = load_workbook(caminho)
    if NOME_SHEET_OBJETOS_ATUALIZADOS not in wb.sheetnames:
        ws = wb.create_sheet(NOME_SHEET_OBJETOS_ATUALIZADOS)
        ws.append(["FUNÇÃO", "OBJETO", "VALOR"])
    else:
        ws = wb[NOME_SHEET_OBJETOS_ATUALIZADOS]

    for objeto, valor in objetos_valores:
        ws.append([role_name, objeto, valor])
    wb.save(caminho)


def _get_session(ambiente: str):
    import win32com.client

    sistema_esperado = MAPA_SISTEMA.get(str(ambiente or "").strip().upper())
    if not sistema_esperado:
        raise ValueError(f"Ambiente inválido: '{ambiente}'. Use DEV, QAD ou PRD.")

    SapGuiAuto = win32com.client.GetObject("SAPGUI")
    application = SapGuiAuto.GetScriptingEngine
    session = next(
        (sess for conn in application.Children for sess in conn.Children
         if sess.Info.SystemName.upper() == sistema_esperado),
        None,
    )
    if not session:
        raise Exception(f"Não encontrei sessão SAP GUI aberta para '{ambiente}' ({sistema_esperado}).")
    return session


def obter_objetos_com_campo(role_name: str, campo: str, ambiente: str = "PRD") -> List[str]:
    """Somente leitura (RFC): objetos de autorização da função que têm o campo indicado."""
    from sap_rfc._rfc_common import (
        build_connection_params_for, load_project_env, find_project_root,
        make_read_only_guard, make_option_eq, read_table,
    )
    from pyrfc import Connection

    campo = campo.strip().upper()
    load_project_env(find_project_root())
    params = build_connection_params_for(ambiente)
    guard = make_read_only_guard(("AGR_1251",))
    conn = Connection(**params)
    try:
        rows = read_table(
            conn, guard,
            table_name="AGR_1251",
            fields=["AGR_NAME", "OBJECT", "FIELD", "DELETED"],
            options=make_option_eq("AGR_NAME", role_name.strip().upper()),
            rowcount=0,
        )
    finally:
        conn.close()

    objetos = sorted({
        obj.strip()
        for _agr, obj, campo_linha, deleted in rows
        if campo_linha.strip().upper() == campo and deleted.strip().upper() != "X"
    })
    return objetos


class EliminarCampoPage:
    def __init__(self, sess):
        self.sess = sess

    def _find(self, sap_id):
        try:
            return self.sess.findById(sap_id)
        except Exception:
            return None

    def _exists(self, sap_id):
        return self._find(sap_id) is not None

    def _busy(self):
        try:
            return bool(getattr(self.sess, "Busy", False))
        except Exception:
            return False

    def _esperar_livre(self, timeout=8.0, pausa=0.05):
        limite = time.time() + timeout
        while time.time() < limite:
            if not self._busy():
                return True
            time.sleep(pausa)
        return False

    def _sbar(self):
        try:
            sbar = self.sess.findById("wnd[0]/sbar")
            return getattr(sbar, "MessageType", "").strip().upper(), (sbar.Text or "").strip()
        except Exception:
            return "", ""

    def _step(self, descricao, path, acao="press", valor=None, vkey=None):
        print(f"  🔎 {descricao}")
        try:
            elem = self.sess.findById(path) if path else None
            if acao == "text":
                elem.text = valor
            elif acao == "press":
                elem.press()
            elif acao == "select":
                if hasattr(elem, "selected"):
                    elem.selected = True
                else:
                    elem.select()
            elif acao == "unselect":
                elem.selected = False
            elif acao == "key":
                elem.key = valor
            elif acao == "sendVKey":
                (elem if path else self.sess.findById("wnd[0]")).sendVKey(vkey)
            self._esperar_livre()
        except Exception as e:
            raise Exception(f"Falha no passo '{descricao}' (ID: {path}): {e}")

    def eliminar_autorizacao_global_campo(self, role_name: str, campo: str) -> Dict[str, Any]:
        campo = campo.strip().upper()
        print(f"\n▶ Função {role_name} — a ELIMINAR autorização global de {campo} (todos os objetos, numa só execução)")

        self._step("Chamar /nPFCGMASSVAL", "wnd[0]/tbar[0]/okcd", "text", "/nPFCGMASSVAL")
        self._step("Enter", "wnd[0]", "sendVKey", vkey=0)

        self._step("Selecionar Execução Direta", "wnd[0]/usr/radMOD_EXE", "select")
        self._step("Selecionar 'Campo para vários objetos'", "wnd[0]/usr/radSEL_FLD", "select")
        self._step("Enter (atualizar ecrã)", "wnd[0]", "sendVKey", vkey=0)

        self._step("Preencher ROLE-LOW", "wnd[0]/usr/ctxtROLE-LOW", "text", role_name)
        self._step("Modo 'D - Eliminar'", "wnd[0]/usr/cmbFLDACT", "key", "D")
        self._step("Objeto em branco (todos os objetos)", "wnd[0]/usr/ctxtFLDOBJ", "text", "")
        self._step(f"Nome do campo = {campo}", "wnd[0]/usr/ctxtFLDFLD", "text", campo)
        self._step(
            "Desmarcar 'Nenhuma mudança para o status Modificado'",
            "wnd[0]/usr/chkS_ACHD", "unselect",
        )
        self._step("Enter (validar seleção)", "wnd[0]", "sendVKey", vkey=0)

        if not self._exists("wnd[0]/usr/btnPFLDN"):
            self._step("Voltar ao início (/N)", "wnd[0]/tbar[0]/okcd", "text", "/N")
            self._step("Enter", "wnd[0]", "sendVKey", vkey=0)
            return {"ok": False, "skip": False, "role": role_name, "campo": campo, "message": "Botão 'Vals.' (valores) não encontrado."}

        self._step("Abrir 'Vals.' (valores a eliminar)", "wnd[0]/usr/btnPFLDN", "press")
        time.sleep(0.4)

        tem_ges1 = self._exists("wnd[1]/usr/btnGES1")
        if not tem_ges1:
            if self.sess.Children.Count > 1:
                self._step("Cancelar popup de valores", "wnd[1]/tbar[0]/btn[12]", "press")
            self._step("Voltar ao início (/N)", "wnd[0]/tbar[0]/okcd", "text", "/N")
            self._step("Enter", "wnd[0]", "sendVKey", vkey=0)
            return {
                "ok": False, "skip": True, "role": role_name, "campo": campo,
                "message": "Botão 'Autorização global para' não apareceu (campo pode não ser de nível organizacional) — a saltar.",
            }

        self._step("Autorização global para (marcar para eliminar)", "wnd[1]/usr/btnGES1", "press")
        time.sleep(0.4)
        if self._exists("wnd[2]/usr/btnBUTTON_1"):
            self._step("Confirmar 'Inserir autorização global' (Sim)", "wnd[2]/usr/btnBUTTON_1", "press")

        self._step("Confirmar popup de valores", "wnd[1]/tbar[0]/btn[0]", "press")

        self._step("Executar (relógio)", "wnd[0]/tbar[1]/btn[8]", "press")
        time.sleep(0.8)

        mt, sb = self._sbar()
        if mt in ("E", "A"):
            return {"ok": False, "skip": False, "role": role_name, "campo": campo, "message": sb or "Erro ao executar."}

        if self._exists("wnd[0]/tbar[1]/btn[20]"):
            self._step("Gerar perfis", "wnd[0]/tbar[1]/btn[20]", "press")
            time.sleep(0.6)
            if self.sess.Children.Count > 1:
                self._step("Fechar popup de geração de perfil", "wnd[1]/tbar[0]/btn[12]", "press")

        self._step("Voltar ao início (/N)", "wnd[0]/tbar[0]/okcd", "text", "/N")
        self._step("Enter", "wnd[0]", "sendVKey", vkey=0)

        return {"ok": True, "skip": False, "role": role_name, "campo": campo, "message": sb or "Autorização global eliminada e perfil gerado."}


def verificar_sem_valor_real_campo(role_name: str, campo: str, ambiente: str = "PRD") -> Dict[str, Any]:
    """Somente leitura (RFC): confirma que nenhum objeto ficou com um valor real
    (qualquer coisa que não seja o placeholder '$<CAMPO>') para este campo."""
    from sap_rfc._rfc_common import (
        build_connection_params_for, load_project_env, find_project_root,
        make_read_only_guard, make_option_eq, read_table,
    )
    from pyrfc import Connection

    role_name = role_name.strip().upper()
    campo = campo.strip().upper()
    placeholder = f"${campo}"
    load_project_env(find_project_root())
    params = build_connection_params_for(ambiente)
    guard = make_read_only_guard(("AGR_1251",))
    conn = Connection(**params)
    try:
        rows = read_table(
            conn, guard,
            table_name="AGR_1251",
            fields=["OBJECT", "AUTH", "FIELD", "LOW", "DELETED"],
            options=make_option_eq("AGR_NAME", role_name),
            rowcount=0,
        )
    finally:
        conn.close()

    valores_por_objeto: Dict[str, set] = {}
    for obj, _auth, campo_linha, low, deleted in rows:
        if deleted.strip().upper() == "X":
            continue
        if campo_linha.strip().upper() != campo:
            continue
        valores_por_objeto.setdefault(obj.strip(), set()).add(low.strip())

    problemas = [
        {"objeto": obj, "valores": sorted(vals)}
        for obj, vals in sorted(valores_por_objeto.items())
        if vals - {placeholder}
    ]
    objetos_ok = [obj for obj, vals in sorted(valores_por_objeto.items()) if not (vals - {placeholder})]

    return {
        "ok": len(problemas) == 0,
        "role": role_name,
        "campo": campo,
        "total_objetos": len(valores_por_objeto),
        "objetos_ok": objetos_ok,
        "problemas": problemas,
    }


def executar_individual(ambiente: str, role_name: str, campo: str) -> Dict[str, Any]:
    role_name = role_name.strip().upper()
    campo = campo.strip().upper()
    session = _get_session(ambiente)

    print(f"A descobrir (RFC, somente leitura) os objetos com {campo} em '{role_name}'...")
    objetos = obter_objetos_com_campo(role_name, campo, ambiente)
    print(f"Objetos encontrados: {len(objetos)}")

    if not objetos:
        return {"ok": True, "skip": True, "role": role_name, "campo": campo, "message": f"Nenhum objeto com {campo} para processar."}

    page = EliminarCampoPage(session)
    try:
        r = page.eliminar_autorizacao_global_campo(role_name, campo)
    except Exception as e:
        r = {"ok": False, "skip": False, "role": role_name, "campo": campo, "message": str(e)}

    if not r["ok"]:
        return r

    print(f"A verificar via RFC (somente leitura) o resultado em '{role_name}'...")
    verificacao = verificar_sem_valor_real_campo(role_name, campo, ambiente)
    r["verificacao"] = verificacao
    if not verificacao["ok"]:
        r["ok"] = False
        r["message"] = (
            f"{r['message']} — mas a verificação RFC encontrou {len(verificacao['problemas'])} "
            f"objeto(s) ainda com valor real de {campo}: {verificacao['problemas']}"
        )
        return r

    registar_objetos_atualizados(role_name, [(obj, "") for obj in verificacao["objetos_ok"]])
    return r


def executar_massa(ambiente: str, roles: List[str], campo: str) -> Dict[str, Any]:
    resultados: List[Dict[str, Any]] = []
    for i, role_name in enumerate(roles, start=1):
        print("\n" + "#" * 75)
        print(f"# [{i}/{len(roles)}] {role_name}")
        print("#" * 75)
        try:
            r = executar_individual(ambiente, role_name, campo)
        except Exception as e:
            r = {"ok": False, "skip": False, "role": role_name, "campo": campo, "message": f"Exceção não tratada: {e}"}

        resultados.append(r)
        if r.get("skip"):
            print(f"   ⏭ Saltado: {r.get('message')}")
            continue
        if not r["ok"]:
            print(f"\n[LOTE ABORTADO] Falha real em '{role_name}': {r.get('message')}")
            break
        print(f"   🟢 {r.get('message')}")

    sucesso = [r for r in resultados if r.get("ok") and not r.get("skip")]
    saltados = [r for r in resultados if r.get("skip")]
    falha = [r for r in resultados if not r.get("ok")]
    restantes = roles[len(resultados):]
    return {
        "ok": len(falha) == 0,
        "resultados": resultados,
        "sucesso": [r["role"] for r in sucesso],
        "saltados": [r["role"] for r in saltados],
        "falha": [r["role"] for r in falha],
        "nao_processados": restantes,
    }


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--ambiente", choices=["DEV", "QAD", "PRD"], required=True)
    parser.add_argument("--role", action="append", default=[], help="Função a processar (repetível para lote).")
    parser.add_argument("--campo", required=True, help="Campo a eliminar (ex.: BUKRS).")
    args = parser.parse_args()

    if len(args.role) <= 1:
        resultado = executar_individual(args.ambiente, args.role[0] if args.role else "Z_PRODUCTION_ORDER_CREATE", args.campo)
    else:
        resultado = executar_massa(args.ambiente, args.role, args.campo)

    print("\n" + "=" * 75)
    print("RESULTADO FINAL")
    print("=" * 75)
    print(resultado)
