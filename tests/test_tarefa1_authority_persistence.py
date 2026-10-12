# -*- coding: utf-8 -*-
"""
tests/test_tarefa1_authority_persistence.py
======================================================================
Testes rigorosos de integração e consistência para a Tarefa 1:
1. Reconciliação PFCG_AUTHORITY [5/8] detecta deltas (novas CREATE e novas COMPOSTA).
2. Deltas de PFCG_AUTHORITY são materializados nas folhas PFCG_CREATE e PFCG_COMPOSTA
   durante a gravação consolidada do Excel.
3. Não duplicação de roles em PFCG_CREATE nem em PFCG_COMPOSTA.
4. Preservação de colunas e dados existentes.
5. Pós-validação com readback confirma delta zero após gravação.
6. Status final reporta pendências em runtime se restarem pendências ou verde se resolvido.
7. Zero escrita SAP (SAP_WRITE = 0).
======================================================================
"""

import os
import shutil
import pytest
import openpyxl
from pathlib import Path
from unittest.mock import patch, MagicMock

import importlib.util

PROJECT_ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location(
    "projeto_perfil",
    PROJECT_ROOT / "Processos" / "Projeto Autorizações" / "Projeto Perfil.py"
)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)


@pytest.fixture
def excel_tarefa1_com_delta_authority(tmp_path):
    """Cria pasta de teste com ficheiro Excel contendo deltas em PFCG_AUTHORITY."""
    p = tmp_path / "S4H_Perfis_teste_tarefa1.xlsx"
    wb = openpyxl.Workbook()

    # CONTROLO
    wb.active.title = "CONTROLO"

    # Proposta
    ws_prop = wb.create_sheet("Proposta")
    ws_prop.append(["FUNÇÃO", "DESCRIÇÃO"])
    ws_prop.append(["Z_ROLE_TESTE", "Role de Teste"])
    ws_prop.append(["VA01", "Criar Encomenda"])

    # Proposta Ativa
    ws_pa = wb.create_sheet("Proposta Ativa")
    ws_pa.append(["UTILIZADOR", "NOME", "COMPOSITE", "DEPARTAMENTO", "DIRECAO", "1", "2", "3", "4", "5", "FUNÇÕES INDIVIDUAIS"])

    # PFCG_CREATE
    ws_cr = wb.create_sheet("PFCG_CREATE")
    ws_cr.append(["ID", "AGR_NAME", "TEXT", "TCODE", "STATUS", "MSG", "TIMESTEMP", "PRD", "QAS", "DEV"])
    ws_cr.append([1, "Z_OLD_SINGLE", "Old Single", "VA01", "Criado", "OK", "", "", "", ""])

    # PFCG_COMPOSTA
    ws_cp = wb.create_sheet("PFCG_COMPOSTA")
    ws_cp.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD", "QAS", "DEV"])
    ws_cp.append([1, "Z_OLD_COMP", "Old Comp", "Z_OLD_SINGLE", "Criado", "OK", "", "", "", ""])

    # PFCG_AUTHORITY com novas roles
    ws_au = wb.create_sheet("PFCG_AUTHORITY")
    ws_au.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "TEXT", "OBJETO", "STATUS", "MSG", "TIMESTEMP"])
    # Nova single Z_NEW_FROM_AUTH sob Z_COMP_AUTH
    ws_au.append([1, "Z_COMP_AUTH", "Composta Authority", "Z_NEW_FROM_AUTH", "Single Nova Auth", "S_TABU_DIS", "", "", ""])
    # Nova single avulsa Z_NEW_SOLO
    ws_au.append([2, "", "", "Z_NEW_SOLO", "Single Nova Solo", "S_TABU_DIS", "", "", ""])
    # Relação já existente (não deve duplicar)
    ws_au.append([3, "Z_OLD_COMP", "Old Comp", "Z_OLD_SINGLE", "Old Single", "S_TABU_DIS", "", "", ""])

    # CUA_ADICIONAR
    ws_cua = wb.create_sheet("CUA_ADICIONAR")
    ws_cua.append(["ID", "UTILIZADOR", "SISTEMA", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP"])

    wb.save(str(p))
    wb.close()
    return str(p)


def test_materializacao_deltas_authority_e_status_final(excel_tarefa1_com_delta_authority, capsys):
    """
    Testa que ao executar a Tarefa 1:
    - PFCG_CREATE recebe Z_NEW_FROM_AUTH e Z_NEW_SOLO.
    - PFCG_COMPOSTA recebe relação (Z_COMP_AUTH, Z_NEW_FROM_AUTH).
    - As linhas já existentes não são duplicadas.
    - Na recarga e pós-validação, o status da etapa 5 fica zerado e o status final é sucesso sem pendências.
    """
    caminho = excel_tarefa1_com_delta_authority

    # Mocks para não rodar SAP RFC real e simular etapas 1, 6 e 7 limpas
    with patch("builtins.input", return_value="N"), \
         patch.object(mod, "auditar_matrizes_departamentais_vs_proposta", return_value={"tem_pendencias": False}), \
         patch.object(mod, "analisar_folha_proposta", return_value={"todos_tcodes": []}), \
         patch.object(mod, "preparar_cua_adicionar", return_value={"novas_linhas_count": 0, "total_users_analisados": 0, "total_atribuicoes_esperadas": 0, "ja_existentes_count": 0, "detalhes_por_dep": {}}), \
         patch("sap_rfc.role_existence_validator_service.executar_etapa_validacao_funcoes_individuais", return_value={"ausentes": 0, "tipos_divergentes": 0}), \
         patch("sap_rfc.role_existence_validator_service.executar_etapa_validacao_funcoes_compostas", return_value={"ausentes": 0, "tipos_divergentes": 0}):

        mod.executar_atualizacao_integrada_excel(
            caminho_excel=caminho,
            validar_tcodes_rfc=False,
            incorporar_prd_rfc=False,
            assumir_sim=True
        )

    out = capsys.readouterr().out
    assert "✓ Sincronização concluída sem pendências." in out

    # Validar persistência física no ficheiro Excel
    wb = openpyxl.load_workbook(caminho)

    # PFCG_CREATE
    ws_cr = wb["PFCG_CREATE"]
    roles_cr = [ws_cr.cell(row=r, column=2).value for r in range(2, ws_cr.max_row + 1)]
    assert "Z_OLD_SINGLE" in roles_cr
    assert "Z_NEW_FROM_AUTH" in roles_cr
    assert "Z_NEW_SOLO" in roles_cr
    # Total de 3 linhas de dados
    assert len(roles_cr) == 3

    # PFCG_COMPOSTA
    ws_cp = wb["PFCG_COMPOSTA"]
    pares_cp = [(ws_cp.cell(row=r, column=2).value, ws_cp.cell(row=r, column=4).value) for r in range(2, ws_cp.max_row + 1)]
    assert ("Z_OLD_COMP", "Z_OLD_SINGLE") in pares_cp
    assert ("Z_COMP_AUTH", "Z_NEW_FROM_AUTH") in pares_cp
    # Total de 2 linhas de dados
    assert len(pares_cp) == 2

    wb.close()


def test_status_final_reporta_pendencias_quando_existirem(excel_tarefa1_com_delta_authority, capsys):
    """
    Testa que se restarem pendências em runtime (ex: TCODE sem mapeamento em [1/8]),
    o status final imprime aviso amarelo com o detalhe.
    """
    caminho = excel_tarefa1_com_delta_authority

    with patch("builtins.input", return_value="1"), \
         patch.object(mod, "auditar_matrizes_departamentais_vs_proposta", return_value={
             "tem_pendencias": True,
             "deps_com_pendencias_count": 1,
             "total_tcodes_sem_mapeamento": 1,
             "pendencias_por_dep": {"IT": [{"tcode": "Z_PEND", "descricao": "Transacao Pendente"}]},
             "tcodes_sem_mapeamento_set": {"Z_PEND"}
         }), \
         patch.object(mod, "analisar_folha_proposta", return_value={"todos_tcodes": []}), \
         patch.object(mod, "preparar_cua_adicionar", return_value={"novas_linhas_count": 0, "total_users_analisados": 0, "total_atribuicoes_esperadas": 0, "ja_existentes_count": 0, "detalhes_por_dep": {}}), \
         patch("sap_rfc.role_existence_validator_service.executar_etapa_validacao_funcoes_individuais", return_value={"ausentes": 0, "tipos_divergentes": 0}), \
         patch("sap_rfc.role_existence_validator_service.executar_etapa_validacao_funcoes_compostas", return_value={"ausentes": 0, "tipos_divergentes": 0}):

        mod.executar_atualizacao_integrada_excel(
            caminho_excel=caminho,
            validar_tcodes_rfc=False,
            incorporar_prd_rfc=False,
            assumir_sim=True
        )

    out = capsys.readouterr().out
    assert "⚠️ Sincronização concluída com pendências conhecidas:" in out
    assert "1 TCODEs sem mapeamento" in out
