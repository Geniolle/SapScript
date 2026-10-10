# -*- coding: utf-8 -*-
"""
tests/test_pfcg_composta_sync_service.py
======================================================================
Bateria de Testes Unitários Aprofundada para a Fase 4
(Sincronização PFCG_COMPOSTA com QAD / PRD e Gestão de Status Excel)
======================================================================
"""

import openpyxl
import pytest
from unittest.mock import MagicMock, patch
from pathlib import Path

from sap_rfc.pfcg_composta_sync_service import (
    obter_pendencias_pfcg_composta,
    validar_roles_filhas_nao_compostas,
    ler_relacoes_agr_agrs,
    sincronizar_pfcg_composta,
    COL_STATUS, COL_MSG, COL_TIMESTEMP, COL_PRD
)


@pytest.fixture
def excel_fake_pfcg_composta(tmp_path):
    """Cria uma planilha Excel temporária simulando a sheet PFCG_COMPOSTA."""
    excel_path = tmp_path / "test_pfcg_composta.xlsx"
    wb = openpyxl.Workbook()
    
    # Sheet default
    ws = wb.active
    ws.title = "PFCG_COMPOSTA"
    
    # Cabeçalho
    ws.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD"])
    
    # Linhas de teste
    # Linha 2: Pendente (STATUS vazio)
    ws.append([742, "Z_BR_PURCHREQ_MANAGER", "PURCHASE REQUISITION MANAGER", "Z_FI_VENDOR_READ", "", "", "", ""])
    # Linha 3: Pendente (STATUS vazio)
    ws.append([743, "Z_BR_PURCHREQ_MANAGER", "PURCHASE REQUISITION MANAGER", "Z_MM_PO_CREATE", "", "", "", ""])
    # Linha 4: Já Concluído (STATUS preenchido)
    ws.append([700, "Z_OLD_ROLE", "OLD DESC", "Z_SINGLE_OLD", "Concluído", "Validado em SAP QAD e PRD via RFC", "2026-10-09 10:00:00", "Validado"])
    
    wb.save(excel_path)
    wb.close()
    return str(excel_path)


def test_obter_pendencias_pfcg_composta(excel_fake_pfcg_composta):
    pendencias = obter_pendencias_pfcg_composta(excel_fake_pfcg_composta)
    assert len(pendencias) == 2
    assert pendencias[0]["id"] == 742
    assert pendencias[0]["composite"] == "Z_BR_PURCHREQ_MANAGER"
    assert pendencias[0]["single"] == "Z_FI_VENDOR_READ"
    assert pendencias[1]["id"] == 743
    assert pendencias[1]["single"] == "Z_MM_PO_CREATE"


def test_validar_roles_filhas_nao_compostas_sucesso():
    mock_conn = MagicMock()
    mock_guard = MagicMock()
    with patch("sap_rfc.pfcg_composta_sync_service.read_table", return_value=[]):
        validas, invalidas = validar_roles_filhas_nao_compostas(mock_conn, mock_guard, ["Z_FI_SINGLE_1", "Z_MM_SINGLE_2"])
        assert validas == {"Z_FI_SINGLE_1", "Z_MM_SINGLE_2"}
        assert len(invalidas) == 0


def test_validar_roles_filhas_nao_compostas_com_composite_invalida():
    mock_conn = MagicMock()
    mock_guard = MagicMock()
    # Simular que Z_COMPOSITE_FILHA está registrada em AGR_FLAGS como COLL_AGR
    mock_data = [["Z_COMPOSITE_FILHA", "COLL_AGR"]]
    with patch("sap_rfc.pfcg_composta_sync_service.read_table", return_value=mock_data):
        validas, invalidas = validar_roles_filhas_nao_compostas(mock_conn, mock_guard, ["Z_FI_SINGLE_1", "Z_COMPOSITE_FILHA"])
        assert validas == {"Z_FI_SINGLE_1"}
        assert invalidas == {"Z_COMPOSITE_FILHA"}


def test_ler_relacoes_agr_agrs():
    mock_conn = MagicMock()
    mock_guard = MagicMock()
    mock_data = [
        ["Z_BR_PURCHREQ_MANAGER", "Z_FI_VENDOR_READ"],
        ["Z_BR_PURCHREQ_MANAGER", "Z_MM_PO_CREATE"]
    ]
    with patch("sap_rfc.pfcg_composta_sync_service.read_table", return_value=mock_data):
        rel, ok, err = ler_relacoes_agr_agrs(mock_conn, mock_guard, ["Z_BR_PURCHREQ_MANAGER"])
        assert ok is True
        assert err is None
        assert "Z_BR_PURCHREQ_MANAGER" in rel
        assert rel["Z_BR_PURCHREQ_MANAGER"] == {"Z_FI_VENDOR_READ", "Z_MM_PO_CREATE"}


def test_sincronizar_pfcg_composta_preflight_falha_qad_mas_prd_ok(excel_fake_pfcg_composta):
    """Testa que se o preflight QAD falhar mas PRD estiver OK, o processamento de PRD avança e gera STATUS='Pendente QAD'."""
    from sap_agent.safety import SafetyGuard
    real_guard = SafetyGuard.build(allow_write_operations=False, allowed_functions=["RFC_READ_TABLE"], allowed_tables=["AGR_DEFINE", "AGR_FLAGS", "AGR_AGRS"])
    mock_conn = MagicMock()
    fake_params = {"user": "U", "passwd": "P", "ashost": "H", "sysnr": "00", "client": "100"}

    with patch("sap_rfc.pfcg_composta_sync_service.preflight_test_connection", side_effect=lambda env: (False, "Falha auth QAD") if env == "QAD" else (True, None)), \
         patch("sap_rfc.pfcg_composta_sync_service.build_connection_params_for_env", return_value=fake_params), \
         patch("sap_rfc.pfcg_composta_sync_service.make_read_only_guard", return_value=real_guard), \
         patch("sap_rfc.pfcg_composta_sync_service.Connection", return_value=mock_conn), \
         patch("sap_rfc.pfcg_composta_sync_service.validar_roles_filhas_nao_compostas", return_value=({"Z_FI_VENDOR_READ", "Z_MM_PO_CREATE"}, set())), \
         patch("sap_rfc.pfcg_composta_sync_service.ler_relacoes_agr_agrs", return_value=({"Z_BR_PURCHREQ_MANAGER": set()}, True, None)):
        
        res = sincronizar_pfcg_composta(excel_fake_pfcg_composta, dry_run=True)
        assert res["status"] == "SUCESSO"
        assert res["detalhes"][0]["status_final"] == "Pendente QAD"


def test_sincronizar_pfcg_composta_preflight_falha_ambos(excel_fake_pfcg_composta):
    """Testa que se o preflight falhar em AMBOS os ambientes, a execução aborta de imediato (PREFLIGHT_FAILED)."""
    with patch("sap_rfc.pfcg_composta_sync_service.preflight_test_connection", return_value=(False, "Ambiente Offline")):
        res = sincronizar_pfcg_composta(excel_fake_pfcg_composta, dry_run=True)
        assert res["status"] == "PREFLIGHT_FAILED"
        assert res["qad_preflight_ok"] is False
        assert res["prd_preflight_ok"] is False


def test_sincronizar_pfcg_composta_dry_run_sucesso(excel_fake_pfcg_composta):
    from sap_agent.safety import SafetyGuard
    real_guard = SafetyGuard.build(allow_write_operations=False, allowed_functions=["RFC_READ_TABLE"], allowed_tables=["AGR_DEFINE", "AGR_FLAGS", "AGR_AGRS"])
    mock_conn = MagicMock()
    fake_params = {"user": "U", "passwd": "P", "ashost": "H", "sysnr": "00", "client": "100"}

    with patch("sap_rfc.pfcg_composta_sync_service.preflight_test_connection", return_value=(True, None)), \
         patch("sap_rfc.pfcg_composta_sync_service.build_connection_params_for_env", return_value=fake_params), \
         patch("sap_rfc.pfcg_composta_sync_service.make_read_only_guard", return_value=real_guard), \
         patch("sap_rfc.pfcg_composta_sync_service.Connection", return_value=mock_conn), \
         patch("sap_rfc.pfcg_composta_sync_service.validar_roles_filhas_nao_compostas", return_value=({"Z_FI_VENDOR_READ", "Z_MM_PO_CREATE"}, set())), \
         patch("sap_rfc.pfcg_composta_sync_service.ler_relacoes_agr_agrs", return_value=({"Z_BR_PURCHREQ_MANAGER": set()}, True, None)):

        res = sincronizar_pfcg_composta(excel_fake_pfcg_composta, dry_run=True)

        assert res["status"] == "SUCESSO"
        assert res["dry_run"] is True
        assert res["total_linhas_pendentes"] == 2
        assert len(res["detalhes"]) == 1
        assert res["detalhes"][0]["status_final"] == "Concluído"

        # Verificar se o Excel PERMANECEU INALTERADO no modo dry-run
        wb = openpyxl.load_workbook(excel_fake_pfcg_composta, data_only=True)
        ws = wb["PFCG_COMPOSTA"]
        assert ws.cell(row=2, column=COL_STATUS).value is None or ws.cell(row=2, column=COL_STATUS).value == ""
        wb.close()


def test_sincronizar_pfcg_composta_sem_pendencias(tmp_path):
    excel_path = tmp_path / "test_empty.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PFCG_COMPOSTA"
    ws.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD"])
    ws.append([700, "Z_ROLE", "TEXT", "Z_SINGLE", "Concluído", "MSG", "TS", "Validado"])
    wb.save(excel_path)
    wb.close()

    res = sincronizar_pfcg_composta(str(excel_path), dry_run=True)
    assert res["status"] == "SEM_PENDENCIAS"
    assert res["total_pendencias"] == 0


def test_sincronizar_pfcg_composta_classificacao_historicas(tmp_path):
    """Valida a separação relacional de Novo Lote (IDs 742-849) e Pendências Históricas (IDs < 742)."""
    excel_path = tmp_path / "test_historico.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PFCG_COMPOSTA"
    ws.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD"])
    ws.append([739, "Z_BR_STOREMAINT", "DESC", "Z_SINGLE_1", "", "", "", ""])
    ws.append([742, "Z_BR_PURCHREQ", "DESC", "Z_SINGLE_2", "", "", "", ""])
    wb.save(excel_path)
    wb.close()

    pend = obter_pendencias_pfcg_composta(str(excel_path))
    assert len(pend) == 2
    assert pend[0]["id"] == 739
    assert pend[0]["categoria"] == "PENDENCIA_HISTORICA"
    assert pend[1]["id"] == 742
    assert pend[1]["categoria"] == "NOVO_LOTE"


def test_integracao_fase3b_para_fase4(tmp_path):
    """Testa a chamada integrada automática da Fase 4 ao final de reconciliar_proposta_ativa_com_prd.py."""
    excel_path = tmp_path / "test_reconciliar.xlsx"
    wb = openpyxl.Workbook()
    ws_ctrl = wb.active
    ws_ctrl.title = "CONTROLO"
    ws_ctrl.append(["DEPARTAMENTO", "STATUS"])
    
    ws_pfcg = wb.create_sheet("PFCG_COMPOSTA")
    ws_pfcg.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD"])

    ws_def = wb.create_sheet("DEFINIÇÕES")
    ws_def.append(["DEPARTAMENTO", "ROLE"])

    ws_prop = wb.create_sheet("Proposta")
    ws_prop.append(["FUNCAO", "DESCRICAO"])

    ws_prop_at = wb.create_sheet("Proposta Ativa")
    ws_prop_at.append(["USER_ID", "COMPOSITE"])

    wb.save(excel_path)
    wb.close()

    from scripts.reconciliar_proposta_ativa_com_prd import executar_reconciliacao_integrada
    res = executar_reconciliacao_integrada(str(excel_path), simular=True)
    assert "fase4_resultado" in res
    assert res["fase4_resultado"]["status"] in ("SEM_PENDENCIAS", "PREFLIGHT_FAILED", "SUCESSO")

