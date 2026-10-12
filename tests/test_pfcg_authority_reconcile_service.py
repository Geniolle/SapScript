# -*- coding: utf-8 -*-
"""
tests/test_pfcg_authority_reconcile_service.py
======================================================================
Testes unitários rigorosos para a reconciliação PFCG_AUTHORITY com
as sheets PFCG_CREATE e PFCG_COMPOSTA.
======================================================================
Cobre todos os cenários mandatados:
1. AGR_NAME_COMPOSTA preenchido => Composite.
2. AGR_NAME => Single.
3. AGR_NAME_COMPOSTA vazio => apenas PFCG_CREATE.
4. Composite + child => cria relação PFCG_COMPOSTA.
5. Relação já existente => não duplica.
6. AGR_NAME repetido => não duplica PFCG_CREATE.
7. Composite repetida com childs diferentes => cria uma relação por child.
8. Composite sem child => não cria relação (AUTHORITY_COMPOSITE_SEM_CHILD).
9. Descrição existente em PFCG_CREATE não é sobrescrita.
10. TCODE existente não é sobrescrito.
11. Nova Single sem TCODE é criada com TCODE vazio.
12. Conflito role como Composite e Single => CONFLITO_TIPO_AUTHORITY.
13. dry_run => EXCEL_WRITE=0.
14. dry_run => SAP_WRITE=0.
15. Modo aplicação cria backup antes de escrever.
16. Falha no backup aborta escrita.
17. Regras das 3 Business Partner deixam de classificá-las incorretamente como Single quando aparecem como AGR_NAME_COMPOSTA.
18. Validação SAP posterior recebe catálogo já reconciliado.
======================================================================
"""

import pytest
import openpyxl
from unittest.mock import patch, MagicMock
from pathlib import Path

from sap_rfc.pfcg_authority_reconcile_service import (
    analisar_reconciliacao_pfcg_authority,
    reconciliar_pfcg_authority_com_catalogos,
    gerar_backup_timestampado,
)
from sap_rfc.role_existence_validator_service import extrair_roles_necessarias_projeto


@pytest.fixture
def temp_excel_cenarios(tmp_path):
    """Cria um ficheiro Excel de teste com diferentes combinações de PFCG_AUTHORITY, PFCG_CREATE e PFCG_COMPOSTA."""
    p = tmp_path / "teste_perfis_authority.xlsx"
    wb = openpyxl.Workbook()

    # Sheet PFCG_CREATE
    ws_cr = wb.active
    ws_cr.title = "PFCG_CREATE"
    ws_cr.append(["ID", "AGR_NAME", "TEXT", "TCODE", "STATUS", "MSG", "TIMESTEMP", "PRD", "QAS", "DEV"])
    ws_cr.append([1, "Z_EXISTING_SINGLE", "Descricao Existente", "VA01", "Criado", "OK", "", "Validado", "", ""])

    # Sheet PFCG_COMPOSTA
    ws_cp = wb.create_sheet("PFCG_COMPOSTA")
    ws_cp.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD", "QAS", "DEV"])
    ws_cp.append([1, "Z_COMPOSITE_A", "Composite A", "Z_EXISTING_SINGLE", "Concluído", "OK", "", "Validado", "", ""])

    # Sheet PFCG_AUTHORITY
    ws_au = wb.create_sheet("PFCG_AUTHORITY")
    ws_au.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "TEXT", "OBJETO", "STATUS", "MSG", "TIMESTEMP"])
    # 1 & 4. Composite + child
    ws_au.append([1, "Z_COMPOSITE_A", "Composite A", "Z_NEW_CHILD_1", "Nova Child 1", "S_TABU_DIS", "", "", ""])
    # 5. Relação já existente (Z_COMPOSITE_A -> Z_EXISTING_SINGLE)
    ws_au.append([2, "Z_COMPOSITE_A", "Composite A", "Z_EXISTING_SINGLE", "Desc Original", "S_TABU_DIS", "", "", ""])
    # 3. AGR_NAME_COMPOSTA vazio => apenas PFCG_CREATE
    ws_au.append([3, "", "", "Z_SINGLE_AVULSA", "Single Avulsa", "S_TABU_DIS", "", "", ""])
    # 6. AGR_NAME repetido
    ws_au.append([4, "", "", "Z_SINGLE_AVULSA", "Single Avulsa Repetida", "S_TABU_DIS", "", "", ""])
    # 7. Composite repetida com childs diferentes
    ws_au.append([5, "Z_COMPOSITE_B", "Composite B", "Z_CHILD_B1", "Child B1", "S_TABU_DIS", "", "", ""])
    ws_au.append([6, "Z_COMPOSITE_B", "Composite B", "Z_CHILD_B2", "Child B2", "S_TABU_DIS", "", "", ""])
    # 8. Composite sem child
    ws_au.append([7, "Z_COMPOSITE_SEM_CHILD", "Composite Sem Child", "", "", "S_TABU_DIS", "", "", ""])

    wb.save(str(p))
    wb.close()
    return str(p)


def test_reconciliacao_cenarios_basicos(temp_excel_cenarios):
    """Testa regras 1, 2, 3, 4, 5, 6, 7, 8."""
    res = analisar_reconciliacao_pfcg_authority(temp_excel_cenarios)

    # 1. AGR_NAME_COMPOSTA preenchido => Composite
    assert "Z_COMPOSITE_A" in res["composites_distintas"]
    assert "Z_COMPOSITE_B" in res["composites_distintas"]
    assert "Z_COMPOSITE_SEM_CHILD" in res["composites_distintas"]

    # 2. AGR_NAME => Single
    assert "Z_NEW_CHILD_1" in res["singles_distintas"]
    assert "Z_SINGLE_AVULSA" in res["singles_distintas"]
    assert "Z_CHILD_B1" in res["singles_distintas"]
    assert "Z_CHILD_B2" in res["singles_distintas"]

    # 3. AGR_NAME_COMPOSTA vazio => entra em novas de PFCG_CREATE mas nenhuma relação em PFCG_COMPOSTA
    novas_create_roles = {x["role"] for x in res["create_novas_necessarias"]}
    assert "Z_SINGLE_AVULSA" in novas_create_roles
    assert not any(x["parent"] == "" for x in res["composta_novas_necessarias"])

    # 4 & 7. Cria relações necessárias em PFCG_COMPOSTA
    novas_compostas = {(x["parent"], x["child"]) for x in res["composta_novas_necessarias"]}
    assert ("Z_COMPOSITE_A", "Z_NEW_CHILD_1") in novas_compostas
    assert ("Z_COMPOSITE_B", "Z_CHILD_B1") in novas_compostas
    assert ("Z_COMPOSITE_B", "Z_CHILD_B2") in novas_compostas

    # 5. Relação já existente não é duplicada
    assert ("Z_COMPOSITE_A", "Z_EXISTING_SINGLE") not in novas_compostas

    # 6. AGR_NAME repetido não gera duplicado
    assert len([x for x in res["create_novas_necessarias"] if x["role"] == "Z_SINGLE_AVULSA"]) == 1

    # 8. Composite sem child não cria relação mas é contada
    assert res["composite_sem_child_count"] == 1
    assert not any(x["parent"] == "Z_COMPOSITE_SEM_CHILD" for x in res["composta_novas_necessarias"])


def test_conflito_role_como_composite_e_single(tmp_path):
    """Cenário 12: Role que aparece como Composite e Single deve ser sinalizada como CONFLITO_TIPO_AUTHORITY."""
    p = tmp_path / "conflito.xlsx"
    wb = openpyxl.Workbook()
    ws_au = wb.active
    ws_au.title = "PFCG_AUTHORITY"
    ws_au.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "TEXT"])
    ws_au.append([1, "Z_ROLE_CONFLITO", "Texto Comp", "Z_CHILD_X", "Texto Child"])
    ws_au.append([2, "Z_OTHER_COMP", "Texto Other", "Z_ROLE_CONFLITO", "Texto Conflito"])
    wb.save(str(p))
    wb.close()

    res = analisar_reconciliacao_pfcg_authority(str(p))
    assert res["conflitos_count"] == 1
    conf = res["conflitos"][0]
    assert conf["role"] == "Z_ROLE_CONFLITO"
    assert conf["status"] == "CONFLITO_TIPO_AUTHORITY"


def test_dry_run_zero_escrita(temp_excel_cenarios):
    """Cenários 13 e 14: dry_run=True => SAP_WRITE=0 e EXCEL_WRITE=0."""
    res = reconciliar_pfcg_authority_com_catalogos(temp_excel_cenarios, dry_run=True, imprimir_log=False)
    assert res["dry_run"] is True
    assert res["sap_write"] == 0
    assert res["excel_write"] == 0
    assert res["backup_path"] is None


def test_modo_aplicacao_com_backup_e_escrita(temp_excel_cenarios):
    """Cenários 9, 10, 11, 15: dry_run=False cria backup, não sobrescreve textos/tcodes existentes, cria novas com tcode vazio."""
    res = reconciliar_pfcg_authority_com_catalogos(temp_excel_cenarios, dry_run=False, imprimir_log=False)
    assert res["ok"] is True
    assert res["dry_run"] is False
    assert res["excel_write"] > 0
    assert res["backup_path"] is not None
    assert Path(res["backup_path"]).exists()

    # Verificar conteúdo gravado no Excel
    wb = openpyxl.load_workbook(temp_excel_cenarios)
    ws_cr = wb["PFCG_CREATE"]

    # 9 & 10. Z_EXISTING_SINGLE preservou TEXT e TCODE
    rows_cr = list(ws_cr.iter_rows(values_only=True))[1:]
    existing_row = next(r for r in rows_cr if r[1] == "Z_EXISTING_SINGLE")
    assert existing_row[2] == "Descricao Existente"
    assert existing_row[3] == "VA01"

    # 11. Nova Single Z_NEW_CHILD_1 foi criada com TCODE vazio
    new_row = next(r for r in rows_cr if r[1] == "Z_NEW_CHILD_1")
    assert new_row[3] == "" or new_row[3] is None

    # Verificar PFCG_COMPOSTA
    ws_cp = wb["PFCG_COMPOSTA"]
    rows_cp = list(ws_cp.iter_rows(values_only=True))[1:]
    pairs = [(r[1], r[3]) for r in rows_cp]
    assert ("Z_COMPOSITE_A", "Z_NEW_CHILD_1") in pairs
    assert ("Z_COMPOSITE_B", "Z_CHILD_B1") in pairs

    wb.close()


def test_falha_no_backup_aborta_escrita(temp_excel_cenarios):
    """Cenário 16: Se o backup falhar, aborta a escrita."""
    with patch("sap_rfc.pfcg_authority_reconcile_service.gerar_backup_timestampado", return_value=(False, "Erro de permissao")):
        res = reconciliar_pfcg_authority_com_catalogos(temp_excel_cenarios, dry_run=False, imprimir_log=False)
        assert res["ok"] is False
        assert "ABORTADO" in res["mensagem"]
        assert res["excel_write"] == 0


def test_business_partner_classificadas_como_composite_em_authority():
    """Cenário 17 e 18: Z_BR_BUSINESS_PARTNER, Z_BR_BUSINESS_PARTNER_GL, Z_BR_TYPE_BP_GERAL reconhecidas como Composite."""
    excel_real = "S4H_Perfis de autorização_v1.xlsx"
    analise = analisar_reconciliacao_pfcg_authority(excel_real)
    bp_info = analise["business_partner_classificacao"]

    for bp in ["Z_BR_BUSINESS_PARTNER", "Z_BR_BUSINESS_PARTNER_GL", "Z_BR_TYPE_BP_GERAL"]:
        assert bp_info[bp]["origem"] == "PFCG_AUTHORITY"
        assert bp_info[bp]["coluna"] == "AGR_NAME_COMPOSTA"
        assert bp_info[bp]["tipo_esperado_resultante"] == "COMPOSITE"

    # E na extração de roles do projeto, aparecem em COMPOSITE e NÃO em SINGLE
    roles_proj = extrair_roles_necessarias_projeto(excel_real)
    for bp in ["Z_BR_BUSINESS_PARTNER", "Z_BR_BUSINESS_PARTNER_GL", "Z_BR_TYPE_BP_GERAL"]:
        assert bp in roles_proj["COMPOSITE"]
        assert bp not in roles_proj["SINGLE"]


def test_formato_log_reconciliacao_compacto_e_com_pendencia(capsys, tmp_path):
    """Testa se o log padrão imprime 1 linha quando não há pendências e exibe detalhes quando há delta."""
    p = tmp_path / "teste_log_compacto.xlsx"
    wb = openpyxl.Workbook()

    # PFCG_CREATE já com todas necessárias
    ws_cr = wb.active
    ws_cr.title = "PFCG_CREATE"
    ws_cr.append(["ID", "AGR_NAME", "TEXT", "TCODE"])
    ws_cr.append([1, "Z_ROLE_S1", "Single 1", ""])
    ws_cr.append([2, "Z_ROLE_S2", "Single 2", ""])

    # PFCG_COMPOSTA já com relação
    ws_cp = wb.create_sheet("PFCG_COMPOSTA")
    ws_cp.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME"])
    ws_cp.append([1, "Z_ROLE_C1", "Comp 1", "Z_ROLE_S1"])

    # PFCG_AUTHORITY sem pendências
    ws_au = wb.create_sheet("PFCG_AUTHORITY")
    ws_au.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "TEXT"])
    ws_au.append([1, "Z_ROLE_C1", "Comp 1", "Z_ROLE_S1", "Single 1"])

    wb.save(str(p))
    wb.close()

    # 1. Sem pendências -> 1 linha
    reconciliar_pfcg_authority_com_catalogos(str(p), dry_run=True, imprimir_log=True)
    captured = capsys.readouterr().out
    assert "[5/8] PFCG_AUTHORITY: 1 linhas | 1C | 1S | CREATE 1/1 | COMPOSTA 1/1 | conflitos 0" in captured
    assert "✓" in captured
    assert "CREATE ausentes:" not in captured
    assert "COMPOSTA ausente:" not in captured

    # 2. Com pendências -> ⚠️ e detalhes
    wb = openpyxl.load_workbook(str(p))
    ws_au = wb["PFCG_AUTHORITY"]
    ws_au.append([2, "Z_ROLE_C2", "Comp 2", "Z_ROLE_S3", "Single 3 (Nova)"])
    wb.save(str(p))
    wb.close()

    reconciliar_pfcg_authority_com_catalogos(str(p), dry_run=True, imprimir_log=True)
    captured_pend = capsys.readouterr().out
    assert "[5/8] PFCG_AUTHORITY:" in captured_pend
    assert "⚠️" in captured_pend
    assert "CREATE ausentes:" in captured_pend
    assert "- Z_ROLE_S3" in captured_pend
    assert "COMPOSTA ausente:" in captured_pend
    assert "- Z_ROLE_C2 -> Z_ROLE_S3" in captured_pend

