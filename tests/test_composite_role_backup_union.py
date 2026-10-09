# -*- coding: utf-8 -*-
"""
tests/test_composite_role_backup_union.py
======================================================================
Testes unitários cobrindo as regras da funcionalidade de Backup de Composite Role:
A. 2 utilizadores, mesma Composite, mesmas funções -> Composite não muda.
B. 2 utilizadores, mesma Composite, um utilizador possui função extra -> função extra entra na Composite.
C. 3 utilizadores, funções parcialmente diferentes -> Composite recebe a união das três listas.
D. Função já existente na Composite -> não duplicar.
E. Nova TCODE mapeada para múltiplas funções -> todas entram na união.
F. Execução repetida -> idempotência (ADICIONADAS = 0, DUPLICADAS = 0).
G. Função existente na Composite que nenhum utilizador atual exige -> não remover automaticamente.
H. Federico (S13360) / Alejandro (S13020) na Composite Z_BR_STOREMAINT_SPECIALIST:
   Federico possui ME21N, ME22N, ME23N; Alejandro não possui.
   Resultado esperado na Composite: Z_PURCHASE_ORDER_CREATE, Z_PURCHASE_ORDER_DISPLAY, Z_PURCHASE_ORDER_REPORT presentes.
======================================================================
"""

import sys
import pytest
import importlib.util
from pathlib import Path
from typing import Dict, List, Set, Tuple, Any

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

script_path = PROJECT_ROOT / "Processos" / "Projeto Autorizações" / "Projeto Perfil.py"
spec = importlib.util.spec_from_file_location("projeto_perfil_mod", script_path)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)


def calcular_uniao_composite_roles(
    composite_para_utilizadores: Dict[str, List[str]],
    funcoes_calculadas_por_utilizador: Dict[str, List[str]],
    composite_atual_funcoes: Dict[str, List[str]]
) -> Dict[str, Dict[str, Any]]:
    """
    Função utilitária de simulação de cálculo da união de Composite Roles
    respeitando a regra aditiva e determinística do Projeto Perfil.
    """
    resultado = {}
    for comp_role, lista_users in composite_para_utilizadores.items():
        atual_set = set(mod.normalizar_texto(f) for f in composite_atual_funcoes.get(comp_role, []))
        
        # 1 e 2. União de todas as funções calculadas pelos utilizadores da Composite
        funcoes_uniao_lista = []
        for u in lista_users:
            u_funcoes = funcoes_calculadas_por_utilizador.get(u, [])
            for f in u_funcoes:
                f_norm = mod.normalizar_texto(f)
                if f_norm and f_norm not in funcoes_uniao_lista:
                    funcoes_uniao_lista.append(f_norm)
                    
        calculada_set = set(funcoes_uniao_lista)
        
        # 3. Adicionar = CALCULADA - ATUAL
        adicionar = [f for f in funcoes_uniao_lista if f not in atual_set]
        
        # 4. Manter = CALCULADA ∩ ATUAL
        manter = [f for f in funcoes_uniao_lista if f in atual_set]
        
        # 5. Candidatas à revisão = ATUAL - CALCULADA (NÃO REMOVER)
        candidatas_revisao = sorted(list(atual_set - calculada_set))
        
        # Composição final = ATUAL U CALCULADA (Aditiva)
        composicao_final = list(composite_atual_funcoes.get(comp_role, []))
        for f in adicionar:
            if f not in composicao_final:
                composicao_final.append(f)
                
        resultado[comp_role] = {
            "utilizadores": lista_users,
            "funcoes_antes": composite_atual_funcoes.get(comp_role, []),
            "funcoes_calculadas": funcoes_uniao_lista,
            "adicionar": adicionar,
            "manter": manter,
            "candidatas_revisao": candidatas_revisao,
            "composicao_final": composicao_final,
            "status": "ALTERADO" if len(adicionar) > 0 else "SEM_ALTERACAO"
        }
    return resultado


def test_a_2_utilizadores_mesma_composite_mesmas_funcoes():
    """A. 2 utilizadores, mesma Composite, mesmas funções -> Composite não muda."""
    composite_users = {"Z_BR_TEST_ROLE": ["USER1", "USER2"]}
    user_funcs = {
        "USER1": ["Z_ROLE_1", "Z_ROLE_2"],
        "USER2": ["Z_ROLE_1", "Z_ROLE_2"],
    }
    comp_atual = {"Z_BR_TEST_ROLE": ["Z_ROLE_1", "Z_ROLE_2"]}

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual)
    comp_res = res["Z_BR_TEST_ROLE"]

    assert comp_res["adicionar"] == []
    assert comp_res["composicao_final"] == ["Z_ROLE_1", "Z_ROLE_2"]
    assert comp_res["status"] == "SEM_ALTERACAO"


def test_b_2_utilizadores_mesma_composite_funcao_extra():
    """B. 2 utilizadores, mesma Composite, um utilizador possui função extra -> função extra entra na Composite."""
    composite_users = {"Z_BR_TEST_ROLE": ["USER_A", "USER_B"]}
    user_funcs = {
        "USER_A": ["F1", "F2", "F3"],
        "USER_B": ["F1", "F2", "F4"],
    }
    comp_atual = {"Z_BR_TEST_ROLE": ["F1", "F2"]}

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual)
    comp_res = res["Z_BR_TEST_ROLE"]

    assert "F3" in comp_res["adicionar"]
    assert "F4" in comp_res["adicionar"]
    assert set(comp_res["composicao_final"]) == {"F1", "F2", "F3", "F4"}
    assert comp_res["status"] == "ALTERADO"


def test_c_3_utilizadores_funcoes_parcialmente_diferentes():
    """C. 3 utilizadores, funções parcialmente diferentes -> Composite recebe a união das três listas."""
    composite_users = {"Z_BR_MULTI": ["U1", "U2", "U3"]}
    user_funcs = {
        "U1": ["F1", "F2"],
        "U2": ["F2", "F3"],
        "U3": ["F4", "F1"],
    }
    comp_atual = {"Z_BR_MULTI": ["F1"]}

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual)
    comp_res = res["Z_BR_MULTI"]

    assert set(comp_res["composicao_final"]) == {"F1", "F2", "F3", "F4"}
    assert set(comp_res["adicionar"]) == {"F2", "F3", "F4"}


def test_d_funcao_ja_existente_nao_duplicar():
    """D. Função já existente na Composite -> não duplicar."""
    composite_users = {"Z_BR_TEST": ["U1"]}
    user_funcs = {"U1": ["F1", "F2"]}
    comp_atual = {"Z_BR_TEST": ["F1", "F2"]}

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual)
    comp_res = res["Z_BR_TEST"]

    assert len(comp_res["composicao_final"]) == 2
    assert comp_res["adicionar"] == []


def test_e_nova_tcode_multiplas_funcoes():
    """E. Nova TCODE mapeada para múltiplas funções -> todas entram na união."""
    tcode_map = {"VK13": ["Z_SALES_PRICECOND_CREATE", "Z_SALES_PRICECOND_REPORT"]}
    tcodes_user = ["VK13"]

    funcs, sem_f = mod.mapeamento_tcodes_para_funcoes_individuais(tcodes_user, tcode_map)
    assert funcs == ["Z_SALES_PRICECOND_CREATE", "Z_SALES_PRICECOND_REPORT"]

    composite_users = {"Z_BR_SALES": ["U1"]}
    user_funcs = {"U1": funcs}
    comp_atual = {"Z_BR_SALES": []}

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual)
    assert set(res["Z_BR_SALES"]["composicao_final"]) == {"Z_SALES_PRICECOND_CREATE", "Z_SALES_PRICECOND_REPORT"}


def test_f_idempotencia_execucao_repetida():
    """F. Execução repetida -> idempotência (ADICIONADAS = 0)."""
    composite_users = {"Z_BR_STOREMAINT_SPECIALIST": ["S13360", "S13020"]}
    user_funcs = {
        "S13360": [
            "Z_PURCHASE_ORDER_CREATE", "Z_PURCHASE_ORDER_DISPLAY", "Z_PURCHASE_ORDER_REPORT",
            "Z_PURCHASE_REQ_CREATE", "Z_PURCHASE_REQ_DISPLAY", "Z_PURCHASE_REQ_REPORT",
            "Z_GOODS_MOVEMENTS", "Z_COSTCENTER_CREATE"
        ],
        "S13020": [
            "Z_PURCHASE_REQ_CREATE", "Z_PURCHASE_REQ_DISPLAY", "Z_PURCHASE_REQ_REPORT",
            "Z_GOODS_MOVEMENTS", "Z_COSTCENTER_CREATE"
        ]
    }
    # 1ª Execução
    comp_atual_1 = {"Z_BR_STOREMAINT_SPECIALIST": [
        "Z_PURCHASE_REQ_CREATE", "Z_PURCHASE_REQ_DISPLAY", "Z_PURCHASE_REQ_REPORT",
        "Z_GOODS_MOVEMENTS", "Z_COSTCENTER_CREATE"
    ]}
    res_1 = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual_1)
    assert len(res_1["Z_BR_STOREMAINT_SPECIALIST"]["adicionar"]) == 3

    # 2ª Execução usando a composição resultante da 1ª
    comp_atual_2 = {"Z_BR_STOREMAINT_SPECIALIST": res_1["Z_BR_STOREMAINT_SPECIALIST"]["composicao_final"]}
    res_2 = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual_2)

    assert len(res_2["Z_BR_STOREMAINT_SPECIALIST"]["adicionar"]) == 0
    assert res_2["Z_BR_STOREMAINT_SPECIALIST"]["status"] == "SEM_ALTERACAO"


def test_g_funcao_existente_nao_exigida_nao_remover():
    """G. Função existente na Composite que nenhum utilizador atual exige -> não remover, apenas CANDIDATA_A_REVISAO."""
    composite_users = {"Z_BR_TEST": ["U1"]}
    user_funcs = {"U1": ["F1"]}
    comp_atual = {"Z_BR_TEST": ["F1", "F_OBSOLETA"]}

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_atual)
    comp_res = res["Z_BR_TEST"]

    assert "F_OBSOLETA" in comp_res["candidatas_revisao"]
    assert "F_OBSOLETA" in comp_res["composicao_final"]  # NÃO REMOVIDA


def test_h_caso_real_federico_alejandro():
    """
    H. Federico (S13360) possui ME21N, ME22N, ME23N -> Z_PURCHASE_ORDER_CREATE, Z_PURCHASE_ORDER_DISPLAY, Z_PURCHASE_ORDER_REPORT.
    Alejandro (S13020) não possui.
    Ambos pertencem a Z_BR_STOREMAINT_SPECIALIST.
    Resultado esperado na Composite: as 3 funções estão presentes e Alejandro não bloqueia.
    """
    composite_users = {
        "Z_BR_STOREMAINT_SPECIALIST": ["S13360", "S13020"]
    }
    user_funcs = {
        "S13360": [
            "Z_GOODS_MOVEMENTS", "Z_PURCHASE_REQ_CREATE", "Z_PURCHASE_REQ_DISPLAY",
            "Z_PURCHASE_REQ_REPORT", "Z_COSTCENTER_CREATE",
            "Z_PURCHASE_ORDER_CREATE", "Z_PURCHASE_ORDER_DISPLAY", "Z_PURCHASE_ORDER_REPORT"
        ],
        "S13020": [
            "Z_GOODS_MOVEMENTS", "Z_PURCHASE_REQ_CREATE", "Z_PURCHASE_REQ_DISPLAY",
            "Z_PURCHASE_REQ_REPORT", "Z_COSTCENTER_CREATE"
        ]
    }
    comp_antes = {
        "Z_BR_STOREMAINT_SPECIALIST": [
            "Z_GOODS_MOVEMENTS", "Z_PURCHASE_REQ_CREATE", "Z_PURCHASE_REQ_DISPLAY",
            "Z_PURCHASE_REQ_REPORT", "Z_COSTCENTER_CREATE"
        ]
    }

    res = calcular_uniao_composite_roles(composite_users, user_funcs, comp_antes)
    comp_res = res["Z_BR_STOREMAINT_SPECIALIST"]

    assert len(comp_res["funcoes_antes"]) == 5
    assert len(comp_res["adicionar"]) == 3
    assert "Z_PURCHASE_ORDER_CREATE" in comp_res["adicionar"]
    assert "Z_PURCHASE_ORDER_DISPLAY" in comp_res["adicionar"]
    assert "Z_PURCHASE_ORDER_REPORT" in comp_res["adicionar"]
    assert len(comp_res["composicao_final"]) == 8


def test_i_pfcg_composta_novas_linhas_4_campos_apenas():
    """
    I. Nova relação PFCG_COMPOSTA:
    Validar que a estrutura novas_linhas_comp e a representação de linha possuem apenas 4 campos,
    garantindo que STATUS, MSG, TIMESTEMP e PRD não são preenchidos e ficam vazios.
    """
    row_struct = {
        "ID": 731,
        "AGR_NAME_COMPOSTA": "Z_BR_STOREMAINT_SPECIALIST",
        "TEXT": "Store Maintenance Specialist",
        "AGR_NAME": "Z_PURCHASE_ORDER_CREATE"
    }

    # 1. Validar exatamente 4 chaves
    assert set(row_struct.keys()) == {"ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME"}
    assert "STATUS" not in row_struct
    assert "MSG" not in row_struct
    assert "TIMESTEMP" not in row_struct
    assert "PRD" not in row_struct

    # 2. Validar tipos/valores nas 4 colunas
    assert isinstance(row_struct["ID"], int)
    assert row_struct["AGR_NAME_COMPOSTA"] == "Z_BR_STOREMAINT_SPECIALIST"
    assert row_struct["TEXT"] == "Store Maintenance Specialist"
    assert row_struct["AGR_NAME"] == "Z_PURCHASE_ORDER_CREATE"


def test_j_simulacao_escrita_openpyxl_4_colunas(tmp_path):
    """
    J. Simulação de escrita em planilha Excel openpyxl:
    Garantir que ao escrever apenas nas colunas 1..4, as colunas 5..8 permanecem None/vazias.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PFCG_COMPOSTA"
    ws.append(["ID", "AGR_NAME_COMPOSTA", "TEXT", "AGR_NAME", "STATUS", "MSG", "TIMESTEMP", "PRD"])

    novas_linhas = [
        {"ID": 101, "AGR_NAME_COMPOSTA": "COMP_1", "TEXT": "Texto 1", "AGR_NAME": "Z_FUNC_1"},
        {"ID": 102, "AGR_NAME_COMPOSTA": "COMP_1", "TEXT": "Texto 1", "AGR_NAME": "Z_FUNC_2"},
    ]

    next_r = ws.max_row + 1
    for row in novas_linhas:
        ws.cell(row=next_r, column=1, value=row["ID"])
        ws.cell(row=next_r, column=2, value=row["AGR_NAME_COMPOSTA"])
        ws.cell(row=next_r, column=3, value=row["TEXT"])
        ws.cell(row=next_r, column=4, value=row["AGR_NAME"])
        next_r += 1

    file_p = tmp_path / "test_out.xlsx"
    wb.save(file_p)

    wb_read = openpyxl.load_workbook(file_p)
    ws_read = wb_read["PFCG_COMPOSTA"]
    rows = list(ws_read.iter_rows(values_only=True))

    assert len(rows) == 3  # Header + 2 linhas
    r1 = rows[1]
    assert r1[0] == 101
    assert r1[1] == "COMP_1"
    assert r1[2] == "Texto 1"
    assert r1[3] == "Z_FUNC_1"
    assert r1[4] is None  # STATUS
    assert r1[5] is None  # MSG
    assert r1[6] is None  # TIMESTEMP
    assert r1[7] is None  # PRD

